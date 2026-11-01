"""从虚构 DOCX 连续生成目录 XMind 与用例 Excel 的完整离线回归。"""

import json
from pathlib import Path
from zipfile import ZipFile

import pymupdf
import pytest
from docx import Document
from openpyxl import load_workbook
from typer.testing import CliRunner

from tc_rag.application.ports.contracts import Capabilities, ModelResult
from tc_rag.application.requests import CasesRequest
from tc_rag.bootstrap import build_pipeline
from tc_rag.core.config import AppConfig
from tc_rag.domain import models as m
from tc_rag.domain.pipeline import Extraction, StageSnapshot
from tc_rag.infrastructure.providers.fakes import FakeEmbedder
from tc_rag.interfaces.workflow import app


class FixtureLLM:
    """按明确测试夹具返回业务对象；没有真实推理能力，不用于演示模型效果。"""

    capabilities = Capabilities(structured_output=True)

    def __init__(self):
        self.calls = 0

    async def generate(self, request, output_schema):
        self.calls += 1
        if request["stage"] == "extract":
            block = request["document_blocks"][0]
            data = {
                "evidence": [
                    {
                        "evidence_id": "e",
                        "block_id": block["block_id"],
                        "locator": block["locator"],
                        "quote": block["content"],
                    }
                ],
                "features": [
                    {
                        "feature_id": "f",
                        "name": "数量校验",
                        "description": "",
                        "evidence_ids": ["e"],
                    }
                ],
                "rules": [
                    {
                        "rule_id": "r",
                        "feature_id": "f",
                        "kind": "field_constraint",
                        "statement": "最多300个，数量为整数。",
                        "evidence_ids": ["e"],
                        "structured_constraint": {
                            "field_path": "数量",
                            "operator": "<=",
                            "boundary": 300,
                            "precision": 0,
                            "unit": "个",
                        },
                    }
                ],
            }
        elif request["stage"] == "plan_points":
            data = {"points": []}
        else:
            rule = request["rules"][0]
            data = {
                "cases": [
                    {
                        "case_id": f"c{i}",
                        "title": p["objective"],
                        "point_ids": [p["point_id"]],
                        "priority": "P2",
                        "steps": [
                            {
                                "action": "执行：" + p["objective"],
                                "expected": "最多300个，数量为整数。",
                                "rule_ids": [rule["rule_id"]],
                                "evidence_ids": rule["evidence_ids"],
                            }
                        ],
                    }
                    for i, p in enumerate(request["points"])
                ]
            }
        return ModelResult(output=output_schema.model_validate(data))


def setup(tmp_path):
    cfg = AppConfig.model_validate(
        {
            "project": {"workdir": str(tmp_path / "work")},
            "storage": {"database": str(tmp_path / "data.sqlite")},
            "chunking": {"token_counter": "character_approx"},
            "planning": {"methods": ["boundary"]},
            "models": {
                "embedding": {
                    "model": "fixture-embedding",
                    "api_key_env": "TC_RAG_EMBEDDING_API_KEY",
                }
            },
            "retrieval": {"embedding_dimension": 3},
        }
    )
    document = tmp_path / "fiction.docx"
    doc = Document()
    doc.add_paragraph("最多300个，数量为整数。")
    doc.save(document)
    config_file = tmp_path / "config.json"
    config_file.write_text(cfg.model_dump_json(), encoding="utf-8")
    model = FixtureLLM()
    pipeline = build_pipeline(cfg, "run_demo", model, FakeEmbedder())
    return pipeline, document, config_file, model


def point_topics(path):
    with ZipFile(path) as archive:
        root = json.loads(archive.read("content.json"))[0]["rootTopic"]

    def walk(node):
        assert "steps" not in node and "expected" not in node
        yield node
        for child in node.get("children", {}).get("attached", []):
            yield from walk(child)

    return [n for n in walk(root) if "point_id" in n.get("tcRag", {})]


@pytest.mark.asyncio
async def test_one_pass_joint_export_and_sqlite_based_regeneration(tmp_path):
    import sqlite3

    pipeline, doc, cfg_file, model = setup(tmp_path)
    result = await pipeline.generate(doc, "生成边界用例", tmp_path / "out", run_id="run_demo")
    assert model.calls == 3
    folder = Path(result["directory"])
    assert folder == tmp_path / "out" / "run_demo"
    assert result["cases"] == 1
    assert Path(result["xmind"]).parent == Path(result["excel"]).parent == folder
    topics = point_topics(Path(result["xmind"]))
    payload = pipeline.read("run_demo", "cases")
    assert {n["tcRag"]["point_id"] for n in topics} == {
        point_id for case in payload["cases"] for point_id in case["point_ids"]
    }
    assert len(topics) == 1
    assert all(case.get("module_path") for case in payload["cases"])
    assert "review_id" not in payload and "reviewed_checksum" not in payload
    book = load_workbook(result["excel"])
    assert book.active.max_column == 11 and book.active.max_row == 2
    assert len(payload["coverage"]) == 1
    assert (folder / "reports" / "quality.md").is_file()
    assert pipeline.read("run_demo", "progress")["current_stage"] == "completed"
    with sqlite3.connect(pipeline.config.storage.database) as connection:
        assert connection.execute("SELECT count(*) FROM review_snapshots").fetchone()[0] == 0
    # XMind 只是导出视图；修改甚至损坏它，不影响从 SQLite 读取事实重新生成。
    Path(result["xmind"]).write_bytes(b"user changed directory view")
    request = CasesRequest.model_validate(pipeline.read("run_demo", "generation_request"))
    repeated = await pipeline.cases(request)
    assert model.calls == 3
    assert repeated["directory"] != result["directory"]
    assert len(point_topics(Path(repeated["xmind"]))) == 1
    runner = CliRunner()
    assert runner.invoke(app, ["status", "run_demo", "--config", str(cfg_file)]).exit_code == 0
    assert (
        runner.invoke(
            app,
            [
                "inspect",
                "run_demo",
                "--stage",
                "cases",
                "--output",
                str(tmp_path / "again" / "cases.xlsx"),
                "--format",
                "xlsx",
                "--config",
                str(cfg_file),
            ],
        ).exit_code
        == 0
    )


@pytest.mark.asyncio
async def test_case_failure_can_retry_saved_points_without_reparsing(tmp_path):
    pipeline, doc, _, _ = setup(tmp_path)

    class FailingOnce(FixtureLLM):
        failed = False

        async def generate(self, request, schema):
            if request["stage"] == "generate" and not self.failed:
                self.failed = True
                raise RuntimeError("fixture-generation-failure")
            return await super().generate(request, schema)

    pipeline.llm = FailingOnce()
    with pytest.raises(RuntimeError, match="fixture-generation-failure"):
        await pipeline.generate(doc, "生成边界用例", tmp_path / "out", run_id="run_demo")
    assert pipeline.read("run_demo", "failure")["stage"] == "generate"
    assert not list((tmp_path / "out").rglob("*.xmind"))
    doc.unlink()  # 证明重试读取已保存事实，不重新访问原文或要求用户回传脑图。
    request = CasesRequest.model_validate(pipeline.read("run_demo", "generation_request"))
    result = await pipeline.cases(request)
    assert result["cases"] == 1
    assert pipeline.read("run_demo", "progress")["current_stage"] == "completed"


@pytest.mark.asyncio
async def test_uncovered_point_triggers_bounded_coverage_repair(tmp_path):
    pipeline, doc, _, _ = setup(tmp_path)

    class OmitFirst(FixtureLLM):
        omitted = False

        async def generate(self, request, schema):
            result = await super().generate(request, schema)
            if request["stage"] == "generate" and not self.omitted:
                self.omitted = True
                result.output.cases = []
            return result

    pipeline.llm = OmitFirst()
    result = await pipeline.generate(doc, "生成边界用例", tmp_path / "out", run_id="run_demo")
    assert result["cases"] == 1
    payload = pipeline.read("run_demo", "cases")
    assert any(
        item["kind"] == "coverage" and "定向补生成" in item["description"]
        for item in payload["issues"]
    )


@pytest.mark.asyncio
async def test_directory_keeps_points_when_generation_is_incomplete(tmp_path):
    pipeline, doc, _, _ = setup(tmp_path)
    pipeline.config.export.formats = ["json"]

    class Incomplete(FixtureLLM):
        async def generate(self, request, schema):
            result = await super().generate(request, schema)
            if request["stage"] == "generate":
                result.output.cases = []
            return result

    pipeline.llm = Incomplete()
    result = await pipeline.generate(doc, "生成边界用例", tmp_path / "out", run_id="run_demo")
    assert result["cases"] == 0 and Path(result["excel"]).is_file()
    topics = point_topics(Path(result["xmind"]))
    assert len(topics) == 1  # XMind 展示全部测试点；Excel 没有可生成用例
    payload = pipeline.read("run_demo", "cases")
    assert len(payload["points"]) == 1
    coverage_issues = [i for i in payload["issues"] if i["kind"] == "coverage"]
    assert len(coverage_issues) == 1
    assert coverage_issues[0]["object_ids"] == [topics[0]["tcRag"]["point_id"]]
    assert pipeline.repository.get("runs", "run_demo", m.Run).status == "partial_failed"


@pytest.mark.asyncio
async def test_empty_generation_exports_empty_directory_and_gap_report(tmp_path):
    pipeline, doc, _, _ = setup(tmp_path)

    class Empty(FixtureLLM):
        async def generate(self, request, schema):
            if request["stage"] == "generate":
                return ModelResult(output=schema())
            return await super().generate(request, schema)

    pipeline.llm = Empty()
    result = await pipeline.generate(doc, "生成用例", tmp_path / "out", run_id="run_demo")
    assert result["cases"] == 0 and result["issues"] > 0
    assert len(point_topics(Path(result["xmind"]))) == 1
    assert load_workbook(result["excel"]).active.max_row == 1


@pytest.mark.asyncio
async def test_failure_checkpoint_and_resume(tmp_path):
    pipeline, doc, _, model = setup(tmp_path)
    pipeline.max_context = 10
    with pytest.raises(ValueError, match="上下文"):
        await pipeline.points(doc, "生成测试点", tmp_path / "out", run_id="run_demo")
    assert model.calls == 0
    assert pipeline.read("run_demo", "failure")["stage"] == "extract"
    pipeline.max_context = 12000
    await pipeline.points(
        doc,
        "生成测试点",
        tmp_path / "out",
        run_id="run_resumed",
        resume="run_demo",
        from_stage="extract",
    )
    assert model.calls == 2
    assert pipeline.repository.get("runs", "run_resumed", m.Run).status != "failed"


@pytest.mark.asyncio
async def test_resume_reuses_only_completed_extraction_batches_from_parent(tmp_path):
    pipeline, doc, _, _ = setup(tmp_path)
    pipeline.config.models.llm.concurrency = 1
    document = Document()
    for title in ("第一模块", "第二模块"):
        document.add_heading(title, 1)
        document.add_paragraph("最多300个，数量为整数。")
    document.save(doc)

    class FailSecondExtraction(FixtureLLM):
        extraction_calls = 0

        async def generate(self, request, schema):
            if request["stage"] == "extract":
                self.extraction_calls += 1
                if self.extraction_calls == 2:
                    raise RuntimeError("fixture-batch-failure")
            return await super().generate(request, schema)

    model = FailSecondExtraction()
    pipeline.llm = model
    with pytest.raises(RuntimeError, match="fixture-batch-failure"):
        await pipeline.points(doc, "生成测试点", tmp_path / "out", run_id="run_parent")
    assert model.extraction_calls == 2
    await pipeline.points(
        doc,
        "生成测试点",
        tmp_path / "out",
        run_id="run_child",
        resume="run_parent",
        from_stage="extract",
    )
    assert model.extraction_calls == 3  # 首批直接复用，只有失败的第二批重新请求。
    assert pipeline.cache_hits == 1
    assert pipeline.read("run_child", "extract")["rules"]


@pytest.mark.asyncio
async def test_pdf_page_checkpoints_are_reused_on_resume(tmp_path):
    """PDF 解析中断后按页恢复：父运行的页检查点复用到子运行，块结果保持一致。"""
    path = tmp_path / "requirements.pdf"
    with pymupdf.open() as document:
        for index in range(3):
            page = document.new_page()
            page.insert_text(
                (40, 40),
                f"第{index + 1}页：批量操作每次最多选择300个对象，数量为整数；"
                "超过300个时拒绝执行。",
                fontsize=11,
            )
        document.save(path)

    cfg = AppConfig.model_validate(
        {
            "project": {"workdir": str(tmp_path / "work")},
            "storage": {"database": str(tmp_path / "data.sqlite")},
            "chunking": {"token_counter": "character_approx"},
            "planning": {"methods": ["boundary"]},
            "models": {
                "embedding": {
                    "model": "fixture-embedding",
                    "api_key_env": "TC_RAG_EMBEDDING_API_KEY",
                }
            },
            "retrieval": {"embedding_dimension": 3},
        }
    )
    model = FixtureLLM()
    pipeline = build_pipeline(cfg, "run_demo", model, FakeEmbedder())

    class FlakyPdf:
        """包装真实解析器，在指定页抛错以模拟解析中途中断。"""

        def __init__(self, inner):
            self.inner = inner
            self.fail_at = 1  # 第二页（0 基）

        def parse(self, asset):
            return self.inner.parse(asset)

        def session(self, asset):
            return _FlakySession(self.inner.session(asset), self)

    class _FlakySession:
        def __init__(self, context, owner):
            self.context = context
            self.owner = owner

        def __enter__(self):
            self.session = self.context.__enter__()
            return self

        def __exit__(self, *exc):
            return self.context.__exit__(*exc)

        @property
        def page_count(self):
            return self.session.page_count

        @property
        def heading_sizes(self):
            return self.session.heading_sizes

        def parse_page(self, index, state):
            if self.owner.fail_at is not None and index == self.owner.fail_at:
                self.owner.fail_at = None
                raise RuntimeError("flaky-parse")
            return self.session.parse_page(index, state)

    from tc_rag.infrastructure.parsing.pdf import PdfParser

    pipeline.parser = lambda suffix: FlakyPdf(PdfParser(pipeline.store))
    with pytest.raises(RuntimeError, match="flaky-parse"):
        await pipeline.points(path, "生成测试点", tmp_path / "out", run_id="run_parent")
    # 中断发生在解析阶段，父运行没有 chunk 快照，但已完成页已有检查点。
    assert pipeline.repository.get("stage_snapshots", "run_parent:chunk", StageSnapshot) is None
    first_page = pipeline.repository.get("stage_snapshots", "run_parent:pdf_page:1", StageSnapshot)
    assert first_page is not None

    pipeline.parser = lambda suffix: PdfParser(pipeline.store)
    await pipeline.points(
        path,
        "生成测试点",
        tmp_path / "out",
        run_id="run_child",
        resume="run_parent",
        from_stage="extract",
    )
    child_blocks = pipeline.read("run_child", "chunk")["blocks"]
    assert len(child_blocks) > 0
    # 已完成页从父运行检查点复用并写入子运行，其余页重新解析。
    child_first = pipeline.repository.get("stage_snapshots", "run_child:pdf_page:1", StageSnapshot)
    child_third = pipeline.repository.get("stage_snapshots", "run_child:pdf_page:3", StageSnapshot)
    assert child_first is not None and child_first.payload == first_page.payload
    assert child_third is not None


@pytest.mark.asyncio
async def test_later_oversized_chunk_is_rejected_before_any_model_call(tmp_path):
    """第一块可发送而后续块超限时，应整组预检失败，避免先产生无效消耗。"""
    from tc_rag.application.ports.errors import ContextBudgetError

    pipeline, doc, _, model = setup(tmp_path)
    document = Document()
    document.add_paragraph("最多300个，数量为整数。")
    document.add_paragraph("原子段落不能静默截断。" * 1000)
    document.save(doc)
    pipeline.config.chunking.target_tokens = 50
    pipeline.config.chunking.overlap_tokens = 0
    pipeline.max_context = 12000

    with pytest.raises(ContextBudgetError) as error:
        await pipeline.points(doc, "生成测试点", tmp_path / "out", run_id="run_precheck")
    assert error.value.estimated_tokens > error.value.limit == 12000
    assert model.calls == 0
    assert pipeline.read("run_precheck", "failure")["stage"] == "extract"
    # 再用同一份已解析事实证明首块本可进入模型，失败确实来自后续块。
    parsed = pipeline.read("run_precheck", "chunk")
    first = parsed["chunks"][0]
    first_blocks = [
        b for b in parsed["blocks"] if b["block_id"] in {s["block_id"] for s in first["sources"]}
    ]
    request = {
        "stage": "extract",
        "system": pipeline.prompts["extract"],
        "instruction": "生成测试点",
        "chunk": first,
        "document_blocks": first_blocks,
        "output_schema": Extraction.model_json_schema(),
    }
    pipeline._check_budget(request)


@pytest.mark.asyncio
async def test_truncated_batch_splits_instead_of_repairing_incomplete_json(tmp_path):
    from tc_rag.application.ports.errors import OutputTruncatedError
    from tc_rag.application.services.design import plan_points

    pipeline, _, _, _ = setup(tmp_path)
    seen = []

    class Truncating(FixtureLLM):
        async def generate(self, request, schema):
            seen.append(len(request["points"]))
            if len(request["points"]) > 1:
                raise OutputTruncatedError("fixture")
            return await super().generate(request, schema)

    pipeline.llm = Truncating()
    pipeline.active_run_id = "run_demo"
    pipeline.repository.save(
        "runs", "run_demo", m.Run(run_id="run_demo", status="running", config_snapshot={})
    )
    extraction = Extraction(
        evidence=[
            m.EvidenceSpan(
                evidence_id="e", block_id="b", locator=m.Locator(), quote="最多300个，数量为整数。"
            )
        ],
        rules=[
            m.RequirementRule(
                rule_id="r",
                feature_id="f",
                kind="field_constraint",
                statement="最多300个",
                evidence_ids=["e"],
                structured_constraint=m.StructuredConstraint(
                    field_path="数量", operator="<=", boundary=300, precision=0
                ),
            )
        ],
    )
    base = plan_points(extraction.rules, ["boundary"])[0]
    points = [
        base.model_copy(update={"point_id": f"p{index}", "objective": f"边界{index}"})
        for index in range(3)
    ]
    await pipeline.retriever.prepare("fiction", extraction.evidence)
    output = await pipeline._generate_batch("生成边界用例", points, extraction)
    assert len(output.cases) == 3
    assert seen == [3, 1, 2, 1, 1]


@pytest.mark.asyncio
async def test_generation_sends_point_limit_and_chapter_option_is_wired(tmp_path):
    pipeline, doc, _, _ = setup(tmp_path)
    seen = []

    class Checking(FixtureLLM):
        async def generate(self, request, schema):
            if request["stage"] == "generate":
                seen.append(request["limits"])
                assert schema.model_json_schema()["properties"]["cases"]["maxItems"] == 6
            return await super().generate(request, schema)

    pipeline.llm = Checking()
    pipeline.config.generation.max_cases_per_point = 6
    pipeline.config.export.point_tree_layout = "chapter"
    result = await pipeline.generate(doc, "生成边界用例", tmp_path / "out", run_id="run_demo")
    assert result["cases"] == 1
    assert seen == [{"max_cases_per_point": 6}]
    with ZipFile(result["xmind"]) as archive:
        assert "未定位章节" in archive.read("content.json").decode("utf-8")


@pytest.mark.asyncio
async def test_extract_chunk_failure_skip_records_issue_and_continues(tmp_path):
    pipeline, doc, _, _ = setup(tmp_path)

    class Failing(FixtureLLM):
        async def generate(self, request, schema):
            if request["stage"] == "extract":
                raise RuntimeError("fixture-chunk-failure")
            return await super().generate(request, schema)

    pipeline.llm = Failing()
    pipeline.config.runtime.chunk_failure_policy = "skip"
    result = await pipeline.generate(doc, "生成用例", tmp_path / "out", run_id="run_skip")
    assert result["cases"] == 0
    issues = pipeline.read("run_skip", "issues")["items"]
    assert any("提炼失败" in item["description"] for item in issues)


@pytest.mark.asyncio
async def test_extract_chunk_failure_fail_policy_stops_run(tmp_path):
    pipeline, doc, _, _ = setup(tmp_path)

    class Failing(FixtureLLM):
        async def generate(self, request, schema):
            if request["stage"] == "extract":
                raise RuntimeError("fixture-chunk-failure")
            return await super().generate(request, schema)

    pipeline.llm = Failing()
    pipeline.config.runtime.chunk_failure_policy = "fail"
    with pytest.raises(RuntimeError, match="fixture-chunk-failure"):
        await pipeline.generate(doc, "生成用例", tmp_path / "out", run_id="run_fail")
