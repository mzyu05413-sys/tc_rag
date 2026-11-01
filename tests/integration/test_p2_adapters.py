"""P2 外部协议与降级测试；只使用内存 HTTP 响应和虚构图片。"""

import asyncio
import io
import json

import httpx
import pytest
from openai import AsyncOpenAI
from PIL import Image

from tc_rag.application.ports.contracts import EmbeddingSpace, RankedCandidate, VisionResult
from tc_rag.application.ports.errors import EmbeddingProviderError
from tc_rag.application.services.chunking import verify_evidence
from tc_rag.application.services.vision import VisionEvidenceService
from tc_rag.core.config import Endpoint, Runtime
from tc_rag.domain.models import DocumentBlock, EvidenceSpan, Locator
from tc_rag.infrastructure.providers.embedding import OpenAIEmbedder
from tc_rag.infrastructure.providers.fakes import FakeEmbedder
from tc_rag.infrastructure.providers.rerank import HttpReranker
from tc_rag.infrastructure.providers.vision import OpenAICompatVision
from tc_rag.infrastructure.retrieval.context import RagContext
from tc_rag.infrastructure.retrieval.index import NumpyIndex
from tc_rag.infrastructure.retrieval.keyword import KeywordIndex, rrf
from tc_rag.infrastructure.storage.files import LocalArtifactStore
from tc_rag.infrastructure.storage.sqlite import SqliteRepository


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "base_url,body,code,attempts",
    [
        (
            "https://open.bigmodel.cn/api/paas/v4/",
            {"error": {"code": "1113", "message": "private-sentinel"}},
            "HTTP_429_BIGMODEL_1113",
            1,
        ),
        (
            "https://fixture/",
            {"error": {"code": "1113", "message": "private-sentinel"}},
            "HTTP_429",
            3,
        ),
        (
            "https://open.bigmodel.cn/api/paas/v4/",
            {"error": {"code": "private-sentinel", "message": "private-sentinel"}},
            "HTTP_429",
            3,
        ),
    ],
)
async def test_embedding_balance_error_is_safe_and_not_retried(
    monkeypatch, base_url, body, code, attempts
):
    async def no_sleep(_):
        pass

    monkeypatch.setattr("tc_rag.infrastructure.providers.embedding.asyncio.sleep", no_sleep)
    calls, records = [], []

    def handler(request):
        calls.append(request)
        return httpx.Response(429, json=body)

    client = AsyncOpenAI(
        api_key="fake",
        base_url=base_url,
        max_retries=0,
        http_client=httpx.AsyncClient(transport=httpx.MockTransport(handler)),
    )
    adapter = OpenAIEmbedder(
        Endpoint(model="m", base_url=base_url), Runtime(), "run", records.append, client
    )
    try:
        with pytest.raises(EmbeddingProviderError) as caught:
            await adapter.embed_batch(["fictional"], EmbeddingSpace(base_url, "m", 2))
        assert caught.value.code == code
        assert len(calls) == len(records) == attempts
        assert all(record.error == code for record in records)
        assert "private-sentinel" not in str(caught.value)
        assert "private-sentinel" not in str(records)
    finally:
        await adapter.close()


@pytest.mark.parametrize("metric", ["cosine", "dot", "euclidean"])
def test_faiss_matches_numpy_with_scope_and_reloading(tmp_path, metric):
    pytest.importorskip("faiss")
    from tc_rag.infrastructure.retrieval.faiss_index import FaissIndex

    space = EmbeddingSpace("fixture", "m", 3, False)
    ids = ["foreign", "a", "b"]
    vectors = [[1, 0, 0], [0.8, 0.3, 0], [0, 1, 0]]
    numpy, faiss = NumpyIndex(metric), FaissIndex(metric)
    for index in (numpy, faiss):
        index.build(ids, vectors, space)
    expected = numpy.search([1, 0, 0], 10, {"a", "b"})
    store = LocalArtifactStore(tmp_path)
    asset = faiss.publish(store)
    faiss.load(store, asset, ids, space, faiss.version)
    faiss.vector_ids = [101, 203, 899]
    actual = faiss.search([1, 0, 0], 10, {"a", "b"})
    assert [h.object_id for h in actual] == [h.object_id for h in expected]
    assert [h.raw_score for h in actual] == pytest.approx([h.raw_score for h in expected], abs=1e-6)
    assert faiss.search([1, 0, 0], 1, set()) == []
    faiss.build([], [], space)
    assert faiss.search([1, 0, 0], 2) == []


def test_keyword_scope_and_rrf_rank_fusion():
    index = KeywordIndex(str.split)
    index.build({"a": "审批 待办", "foreign": "审批 审批 审批", "b": "删除"})
    assert index.search("审批", 1, {"a", "b"}) == ["a"]
    fused = rrf([["a", "b"], ["b"]], 2)
    assert fused[0][0] == "b"
    assert fused[0][1] == pytest.approx(1 / 62 + 1 / 61)


@pytest.mark.asyncio
@pytest.mark.parametrize("failure", [False, True])
async def test_http_rerank_protocol_and_fallback(tmp_path, failure):
    requests, records = [], []

    def handler(request):
        requests.append(json.loads(request.content))
        if failure:
            raise httpx.ReadTimeout("secret response must not leak")
        return httpx.Response(
            200,
            json={
                "results": [
                    {"index": 1, "relevance_score": 0.9},
                    {"index": 0, "relevance_score": 0.1},
                ],
                "usage": {"prompt_tokens": 5, "total_tokens": 5},
            },
        )

    adapter = HttpReranker(
        Endpoint(adapter="http", model="rerank", base_url="https://fixture/rerank"),
        Runtime(),
        "run",
        records.append,
        httpx.AsyncClient(transport=httpx.MockTransport(handler)),
    )
    repo, store = SqliteRepository(tmp_path / "db"), LocalArtifactStore(tmp_path / "files")
    context = RagContext(
        repo,
        store,
        FakeEmbedder(),
        EmbeddingSpace("fixture", "m", 3),
        "cosine",
        2,
        "run",
        reranker=adapter,
        rerank_top_n=2,
    )
    await context.prepare(
        "doc",
        [
            EvidenceSpan(evidence_id=key, block_id=key, quote=key, locator=Locator())
            for key in ["a", "b"]
        ],
    )
    try:
        bundle = await context.retrieve("query", set())
        original = [h["object_id"] for h in bundle["vector_hits"]]
        actual = [h["object_id"] for h in bundle["hits"]]
        assert actual == (original if failure else list(reversed(original)))
        assert bundle["rerank"]["status"] == ("fallback" if failure else "succeeded")
        assert requests[0]["documents"] == original
        assert requests[0]["top_n"] == 2
        assert records[0].stage == "rerank"
        assert "secret" not in json.dumps(bundle)
    finally:
        await adapter.close()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "rows",
    [
        [{"index": 0, "relevance_score": 0.1}, {"index": 0, "relevance_score": 0.2}],
        [{"index": 5, "relevance_score": 0.1}],
        [{"index": 0, "relevance_score": "invalid"}],
    ],
)
async def test_rerank_rejects_malformed_results(rows):
    client = httpx.AsyncClient(
        transport=httpx.MockTransport(lambda _: httpx.Response(200, json={"results": rows}))
    )
    adapter = HttpReranker(
        Endpoint(adapter="http", model="rerank", base_url="https://fixture/rerank"),
        Runtime(),
        "run",
        lambda _: None,
        client,
    )
    try:
        with pytest.raises(RuntimeError):
            await adapter.rank("query", [RankedCandidate("a", "a"), RankedCandidate("b", "b")])
    finally:
        await adapter.close()


@pytest.mark.asyncio
async def test_embedding_retry_budget_and_dimensions(monkeypatch):
    async def no_sleep(_):
        pass

    monkeypatch.setattr("tc_rag.infrastructure.providers.embedding.asyncio.sleep", no_sleep)
    calls, records = [], []

    def handler(request):
        calls.append(request)
        if len(calls) == 1:
            return httpx.Response(
                429, json={"error": {"message": "private"}}, headers={"retry-after": "0"}
            )
        return httpx.Response(
            200,
            json={
                "data": [{"index": 0, "embedding": [3.0, 4.0]}],
                "usage": {"prompt_tokens": 2, "total_tokens": 2},
                "model": "m",
            },
        )

    client = AsyncOpenAI(
        api_key="fake",
        max_retries=0,
        http_client=httpx.AsyncClient(transport=httpx.MockTransport(handler)),
    )
    adapter = OpenAIEmbedder(
        Endpoint(model="m"), Runtime(budget={"max_calls": 2}), "run", records.append, client
    )
    try:
        result = await adapter.embed_batch(["a"], EmbeddingSpace("fixture", "m", 2))
        assert result.vectors == [[0.6, 0.8]]
        assert [r.error for r in records] == ["HTTP_429", None]
        with pytest.raises(RuntimeError, match="预算"):
            await adapter.embed_batch(["a"], EmbeddingSpace("fixture", "m", 2))
        assert len(calls) == 2
    finally:
        await adapter.close()


@pytest.mark.asyncio
async def test_vision_cache_and_provenance(tmp_path):
    class Reader:
        calls = 0

        async def read(self, image, context):
            self.calls += 1
            return VisionResult("最多300个，数量为整数。", "fixture-vision", "vision/v1")

    repo, store = SqliteRepository(tmp_path / "db"), LocalArtifactStore(tmp_path / "files")
    asset = store.put("image.png", b"fixture-only-reader")
    block = DocumentBlock(
        block_id="image",
        document_id="doc",
        kind="image",
        order=0,
        content=asset.model_dump_json(),
        locator=Locator(page=1),
    )
    reader = Reader()
    service = VisionEvidenceService(repo, store, reader, {"model": "fixture", "preprocess": "v1"})
    blocks, issues = await service.process([block])
    again, _ = await service.process([block])
    assert reader.calls == 1 and blocks == again and issues
    derived = blocks[1]
    evidence = EvidenceSpan(
        evidence_id="e",
        block_id=derived.block_id,
        locator=derived.locator,
        quote=derived.content,
        origin=derived.origin,
        processing_model=derived.processing_model,
        processing_version=derived.processing_version,
    )
    assert evidence.uncertain and verify_evidence(evidence, derived)
    assert not verify_evidence(evidence.model_copy(update={"origin": "text"}), derived)
    changed = VisionEvidenceService(repo, store, reader, {"model": "new", "preprocess": "v1"})
    await changed.process([block])
    assert reader.calls == 2


@pytest.mark.asyncio
async def test_vision_processes_images_concurrently_and_preserves_order(tmp_path):
    class SlowReader:
        def __init__(self) -> None:
            self.active = 0
            self.max_active = 0
            self.calls = 0

        async def read(self, image, context):
            self.calls += 1
            self.active += 1
            self.max_active = max(self.max_active, self.active)
            try:
                await asyncio.sleep(0.05)
                if context.endswith("fail"):
                    raise RuntimeError("boom")
                return VisionResult(f"desc:{context}", "fixture-vision", "vision/v1")
            finally:
                self.active -= 1

    repo = SqliteRepository(tmp_path / "db")
    store = LocalArtifactStore(tmp_path / "files")
    blocks: list[DocumentBlock] = []
    for index in range(5):
        asset = store.put(f"image{index}.png", f"bytes-{index}".encode())
        blocks.append(
            DocumentBlock(
                block_id=f"img{index}",
                document_id="doc",
                kind="image",
                order=index,
                heading_path=["fail"] if index == 3 else [f"img{index}"],
                content=asset.model_dump_json(),
                locator=Locator(page=1),
            )
        )
        if index == 2:
            blocks.append(
                DocumentBlock(
                    block_id="para",
                    document_id="doc",
                    kind="paragraph",
                    order=0,
                    content="plain",
                    locator=Locator(page=1),
                )
            )

    reader = SlowReader()
    service = VisionEvidenceService(repo, store, reader, {"model": "fixture"}, concurrency=2)
    result, issues = await service.process(blocks)

    # 多张图片确实并行处理，而不是逐张串行等待。
    assert reader.max_active == 2
    # 每个图片块后紧跟其派生块，普通文字块位置不变。
    ids = [block.block_id for block in result]
    assert ids[0] == "img0"
    assert ids[1].startswith("img0_vision_")
    assert ids.index("img2") < ids.index("para") < ids.index("img3")
    # 失败图片不产生派生块，但原块保留并记录 unparsed 问题；其余仍成功。
    assert not any(block_id.startswith("img3_vision_") for block_id in ids)
    assert ids.count("img3") == 1
    assert any(item.kind == "unparsed" for item in issues)
    assert sum(block.origin == "vision" for block in result) == 4


@pytest.mark.asyncio
async def test_vision_cancellation_drains_workers_and_leaves_queue_unstarted(tmp_path):
    """取消后等待在途任务退出，队列中的图片不再调用模型。"""
    started = asyncio.Event()

    class WaitingReader:
        active = 0
        calls = 0

        async def read(self, image, context):
            self.calls += 1
            self.active += 1
            if self.active == 2:
                started.set()
            try:
                await asyncio.Event().wait()
            finally:
                self.active -= 1

    reader = WaitingReader()
    store = LocalArtifactStore(tmp_path / "files")
    asset = store.put("image.png", b"image")
    blocks = [
        DocumentBlock(
            block_id=f"img{i}",
            document_id="doc",
            kind="image",
            order=i,
            content=asset.model_dump_json(),
            locator=Locator(),
        )
        for i in range(6)
    ]
    service = VisionEvidenceService(
        SqliteRepository(tmp_path / "db"), store, reader, {"model": "fixture"}, concurrency=2
    )
    task = asyncio.create_task(service.process(blocks))
    await asyncio.wait_for(started.wait(), timeout=2)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert reader.active == 0 and reader.calls == 2


@pytest.mark.asyncio
async def test_vision_http_payload_and_capability_gate():
    endpoint = Endpoint(model="vision", supports_images=True)
    records, requests = [], []
    buffer = io.BytesIO()
    Image.new("RGB", (2000, 500), "white").save(buffer, format="PNG")

    def handler(request):
        requests.append(json.loads(request.content))
        return httpx.Response(
            200,
            json={
                "id": "fake",
                "created": 0,
                "model": "vision",
                "object": "chat.completion",
                "choices": [
                    {
                        "index": 0,
                        "finish_reason": "stop",
                        "message": {"role": "assistant", "content": "图中有审批节点"},
                    }
                ],
                "usage": {"prompt_tokens": 3, "completion_tokens": 2, "total_tokens": 5},
            },
        )

    client = AsyncOpenAI(
        api_key="fake",
        max_retries=0,
        http_client=httpx.AsyncClient(transport=httpx.MockTransport(handler)),
    )
    adapter = OpenAICompatVision(endpoint, Runtime(), "run", records.append, "描述", client)
    try:
        result = await adapter.read(buffer.getvalue(), "审批")
        assert result.text == "图中有审批节点"
        assert requests[0]["messages"][1]["content"][1]["image_url"]["url"].startswith(
            "data:image/png;"
        )
        assert records[0].usage.total_tokens == 5
        with pytest.raises(ValueError, match="supports_images"):
            OpenAICompatVision(
                endpoint.model_copy(update={"supports_images": False}),
                Runtime(),
                "run",
                records.append,
                "描述",
                client,
            )
    finally:
        await adapter.close()


def test_vector_ids_are_persistent_int64(tmp_path):
    repo = SqliteRepository(tmp_path / "db")
    first = repo.allocate_vector_ids("space1", ["a", "b"])
    reopened = SqliteRepository(tmp_path / "db")
    assert reopened.allocate_vector_ids("space1", ["b", "a"]) == list(reversed(first))
    assert reopened.allocate_vector_ids("space2", ["a"])[0] not in first
    assert all(0 < i < 2**63 for i in first)


def test_faiss_cosine_tiny_vectors_and_invalid_query_shape():
    pytest.importorskip("faiss")
    from tc_rag.infrastructure.retrieval.faiss_index import FaissIndex

    index = FaissIndex("cosine")
    index.build(["a", "b"], [[1e-30, 0], [0, 1e-30]], EmbeddingSpace("fixture", "m", 2, False))
    hits = index.search([1e-30, 0], 2)
    assert hits[0].object_id == "a" and hits[0].raw_score == pytest.approx(1)
    with pytest.raises(ValueError, match="一维"):
        index.search([[1, 0]], 2)


@pytest.mark.asyncio
async def test_failed_publication_preserves_active_generation(tmp_path, monkeypatch):
    repo, store = SqliteRepository(tmp_path / "db"), LocalArtifactStore(tmp_path / "files")
    context = RagContext(
        repo, store, FakeEmbedder(), EmbeddingSpace("fixture", "m", 3), "cosine", 2, "run"
    )

    def span(key):
        return EvidenceSpan(evidence_id=key, block_id=key, quote=key, locator=Locator())

    await context.prepare("doc-a", [span("a")])
    version, scope = context.index.version, context.scope

    def failed_put(*args):
        raise OSError("simulated disk failure")

    monkeypatch.setattr(store, "put", failed_put)
    with pytest.raises(OSError):
        await context.prepare("doc-b", [span("b")])
    assert context.index.version == version and context.scope == scope
    bundle = await context.retrieve("query", {"a"})
    assert bundle["hits"][0]["object_id"] == "a"


@pytest.mark.asyncio
async def test_cache_disabled_and_corrupted_cache(tmp_path):
    from tc_rag.domain.rag import RagRecord

    repo, store = SqliteRepository(tmp_path / "db"), LocalArtifactStore(tmp_path / "files")
    embedder = FakeEmbedder()
    context = RagContext(
        repo,
        store,
        embedder,
        EmbeddingSpace("fixture", "m", 3),
        "cosine",
        2,
        "run",
        cache_enabled=False,
    )
    await context.vector("query")
    await context.vector("query")
    assert embedder.calls == 2
    context.cache_enabled = True
    vector, _ = await context.vector("query")
    from tc_rag.domain.identity import fingerprint

    key = "embedding_" + fingerprint([context.space_id, "query"])
    saved = repo.get("embedding_records", key, RagRecord)
    saved.payload["vector"] = [v + 1 for v in vector]
    repo.save("embedding_records", key, saved)
    with pytest.raises(ValueError, match="校验"):
        await context.vector("query")


@pytest.mark.asyncio
async def test_llm_and_embedding_share_request_budget(tmp_path):
    from tc_rag.bootstrap import build_pipeline
    from tc_rag.core.config import AppConfig
    from tc_rag.domain.pipeline import PointBatch
    from tc_rag.infrastructure.providers.openai_sdk import OpenAIStructuredLLM

    calls = []

    def handler(request):
        calls.append(request)
        return httpx.Response(
            200,
            json={
                "id": "fake",
                "created": 0,
                "object": "chat.completion",
                "model": "m",
                "choices": [
                    {
                        "index": 0,
                        "finish_reason": "stop",
                        "message": {"role": "assistant", "content": '{"points":[]} '},
                    }
                ],
            },
        )

    cfg = AppConfig.model_validate(
        {
            "project": {"workdir": str(tmp_path / "files")},
            "storage": {"database": str(tmp_path / "db")},
            "models": {"llm": {"model": "m"}, "embedding": {"model": "m"}},
            "retrieval": {"embedding_dimension": 3},
            "runtime": {"budget": {"max_calls": 1}},
        }
    )

    def client():
        return AsyncOpenAI(
            api_key="fake",
            max_retries=0,
            http_client=httpx.AsyncClient(transport=httpx.MockTransport(handler)),
        )

    llm = OpenAIStructuredLLM(cfg.models.llm, cfg.runtime, "run", lambda _: None, client=client())
    embedding = OpenAIEmbedder(cfg.models.embedding, cfg.runtime, "run", lambda _: None, client())
    try:
        build_pipeline(cfg, "run", llm, embedding)
        await llm.generate({"stage": "plan_points", "system": "fixture"}, PointBatch)
        with pytest.raises(RuntimeError, match="预算"):
            await embedding.embed_batch(["query"], EmbeddingSpace("fixture", "m", 3))
        assert len(calls) == 1
    finally:
        await llm.close()
        await embedding.close()


def test_bad_embedding_config_fails_before_creating_llm_client(tmp_path, monkeypatch):
    from tc_rag.bootstrap import build_pipeline
    from tc_rag.core.config import AppConfig

    monkeypatch.setenv("TC_RAG_LLM_API_KEY", "fake")

    def should_not_construct(*args, **kwargs):
        pytest.fail("配置检查失败前不应创建网络客户端")

    monkeypatch.setattr(
        "tc_rag.infrastructure.providers.openai_sdk.AsyncOpenAI", should_not_construct
    )
    cfg = AppConfig.model_validate(
        {
            "project": {"workdir": str(tmp_path / "files")},
            "storage": {"database": str(tmp_path / "db")},
            "models": {"llm": {"model": "m", "api_key_env": "TC_RAG_LLM_API_KEY"}},
        }
    )
    with pytest.raises(ValueError, match="模型端点"):
        build_pipeline(cfg, "run")


@pytest.mark.asyncio
async def test_vision_block_is_preserved_through_extraction_and_retrieval(tmp_path):
    from docx import Document
    from test_pipeline import setup

    class Reader:
        async def read(self, image, context):
            return VisionResult("最多300个，数量为整数。", "fixture-vision", "vision/v1")

    pipeline, document, _, _ = setup(tmp_path)
    buffer = io.BytesIO()
    Image.new("RGB", (20, 20), "white").save(buffer, format="PNG")
    buffer.seek(0)
    doc = Document()
    doc.add_picture(buffer)
    doc.save(document)
    pipeline.config.parsing.vision_enabled = True
    pipeline.vision = VisionEvidenceService(pipeline.repository, pipeline.store, Reader(), {"v": 1})
    pipeline.retriever = RagContext(
        pipeline.repository,
        pipeline.store,
        FakeEmbedder(),
        EmbeddingSpace("fixture", "m", 3),
        "cosine",
        20,
        "run_vision",
    )
    await pipeline.points(document, "生成测试点", tmp_path / "out", run_id="run_vision")
    # 图片解释独立成阶段：解析与视觉各自记录耗时，便于核对控制台与 manifest。
    assert "parse" in pipeline.stages and "vision" in pipeline.stages
    evidence = pipeline.read("run_vision", "extract")["evidence"]
    assert evidence and all(e["origin"] == "vision" and e["uncertain"] for e in evidence)
    assert all(e["processing_model"] == "fixture-vision" for e in evidence)
    supplement = pipeline.read("run_vision", "retrieval")["bundles"][0]["supplemental_evidence"]
    assert supplement and all(e["origin"] == "vision" for e in supplement)


@pytest.mark.asyncio
@pytest.mark.parametrize("vector", [[1.0], [float("nan"), 1.0], [0.0, 0.0]])
async def test_real_embedding_rejects_bad_dimension_nonfinite_or_zero(vector):
    content = json.dumps(
        {
            "model": "m",
            "data": [{"index": 0, "embedding": vector}],
            "usage": {"prompt_tokens": 1, "total_tokens": 1},
        }
    ).encode()
    client = AsyncOpenAI(
        api_key="fake",
        max_retries=0,
        http_client=httpx.AsyncClient(
            transport=httpx.MockTransport(
                lambda _: httpx.Response(
                    200, content=content, headers={"content-type": "application/json"}
                )
            )
        ),
    )
    records = []
    adapter = OpenAIEmbedder(Endpoint(model="m"), Runtime(), "run", records.append, client)
    try:
        with pytest.raises(ValueError):
            await adapter.embed_batch(["text"], EmbeddingSpace("fixture", "m", 2))
        assert len(records) == 1 and records[0].error == "ValueError"
    finally:
        await adapter.close()


@pytest.mark.asyncio
async def test_partial_embedding_failure_does_not_publish_index(tmp_path):
    import sqlite3
    from contextlib import closing

    records, calls = [], []

    def handler(request):
        calls.append(request)
        if len(calls) == 2:
            return httpx.Response(400, json={"error": {"message": "private"}})
        return httpx.Response(
            200,
            json={
                "model": "m",
                "data": [{"index": 0, "embedding": [1.0, 2.0]}],
                "usage": {"prompt_tokens": 1, "total_tokens": 1},
            },
        )

    client = AsyncOpenAI(
        api_key="fake",
        max_retries=0,
        http_client=httpx.AsyncClient(transport=httpx.MockTransport(handler)),
    )
    adapter = OpenAIEmbedder(
        Endpoint(model="m", batch_size=1), Runtime(), "run", records.append, client
    )
    repo, store = SqliteRepository(tmp_path / "db"), LocalArtifactStore(tmp_path / "files")
    context = RagContext(
        repo, store, adapter, EmbeddingSpace("fixture", "m", 2), "cosine", 2, "run"
    )
    spans = [
        EvidenceSpan(evidence_id=key, block_id=key, quote=key, locator=Locator())
        for key in ("a", "b")
    ]
    try:
        with pytest.raises(RuntimeError, match="HTTP_400"):
            await context.prepare("doc", spans)
        assert [r.error for r in records] == [None, "HTTP_400"]
        assert context.index.version == "unbuilt"
        with closing(sqlite3.connect(repo.database)) as connection:
            assert connection.execute("SELECT count(*) FROM index_manifests").fetchone()[0] == 0
    finally:
        await adapter.close()


@pytest.mark.asyncio
async def test_optional_chinese_keyword_rescues_vector_miss(tmp_path):
    pytest.importorskip("jieba")
    repo, store = SqliteRepository(tmp_path / "db"), LocalArtifactStore(tmp_path / "files")
    # 明确构造向量漏召回，验证关键词能把同范围中文证据带回上下文。
    embedder = FakeEmbedder([[[0.0, 1.0], [1.0, 0.0]], [[1.0, 0.0]]])
    context = RagContext(
        repo,
        store,
        embedder,
        EmbeddingSpace("fixture", "m", 2),
        "cosine",
        1,
        "run",
        keyword_enabled=True,
    )
    await context.prepare(
        "doc",
        [
            EvidenceSpan(evidence_id="a", block_id="a", quote="审批待办", locator=Locator()),
            EvidenceSpan(evidence_id="b", block_id="b", quote="批量删除", locator=Locator()),
        ],
    )
    bundle = await context.retrieve("审批", set())
    assert bundle["vector_hits"][0]["object_id"] == "b"
    assert bundle["keyword_ids"] == ["a"]
    assert bundle["hits"][0]["object_id"] == "a"
    assert bundle["hits"][0]["raw_score"] is None
    assert bundle["supplemental_evidence"][0]["quote"] == "审批待办"


@pytest.mark.asyncio
async def test_ui_display_rules_merge_to_one_summary_point(tmp_path):
    from docx import Document
    from test_pipeline import FixtureLLM, setup

    from tc_rag.application.ports.contracts import ModelResult
    from tc_rag.domain.review import PointTree

    class Reader:
        async def read(self, image, context):
            return VisionResult(
                "产品列表界面展示字段A、字段B、列C、列D。",
                "fixture-vision",
                "vision/v1",
            )

    class DisplayLLM(FixtureLLM):
        async def generate(self, request, output_schema):
            if request["stage"] == "extract":
                block = request["document_blocks"][0]
                data = {
                    "evidence": [
                        {
                            "evidence_id": "e",
                            "block_id": block["block_id"],
                            "locator": block["locator"],
                            "quote": block["content"],
                        }
                    ],
                    "features": [
                        {
                            "feature_id": "f",
                            "name": "产品列表",
                            "description": "",
                            "evidence_ids": ["e"],
                        }
                    ],
                    "rules": [
                        {
                            "rule_id": "r1",
                            "feature_id": "f",
                            "kind": "other",
                            "statement": "产品列表展示字段A、字段B",
                            "evidence_ids": ["e"],
                            "status": "needs_review",
                        },
                        {
                            "rule_id": "r2",
                            "feature_id": "f",
                            "kind": "other",
                            "statement": "产品列表展示列C、列D",
                            "evidence_ids": ["e"],
                            "status": "needs_review",
                        },
                    ],
                }
                return ModelResult(output=output_schema.model_validate(data))
            return await super().generate(request, output_schema)

    pipeline, document, _, _ = setup(tmp_path)
    buffer = io.BytesIO()
    Image.new("RGB", (20, 20), "white").save(buffer, format="PNG")
    buffer.seek(0)
    doc = Document()
    doc.add_picture(buffer)
    doc.save(document)
    pipeline.config.parsing.vision_enabled = True
    pipeline.llm = DisplayLLM()
    pipeline.vision = VisionEvidenceService(
        pipeline.repository, pipeline.store, Reader(), {"v": 1}
    )
    pipeline.retriever = RagContext(
        pipeline.repository,
        pipeline.store,
        FakeEmbedder(),
        EmbeddingSpace("fixture", "m", 3),
        "cosine",
        20,
        "run_display",
    )
    await pipeline.points(
        document, "生成测试点", tmp_path / "out", run_id="run_display"
    )
    planned = pipeline.read("run_display", "plan_points")
    assert len(planned["points"]) == 1
    summary = planned["points"][0]
    assert summary["technique"] == "other"
    assert summary["objective"] == "产品列表：界面展示样式与内容正确"
    tree = pipeline.repository.get("point_trees", planned["tree_id"], PointTree)
    assert tree is not None
    assert len([node for node in tree.nodes if node.kind == "point"]) == 1
    issues = pipeline.read("run_display", "issues")["items"]
    assert any("归并" in item["description"] for item in issues)
    manifest = pipeline.read("run_display", "manifest")
    assert manifest["prompts"]["plan_points"]["version"] == "v2"
