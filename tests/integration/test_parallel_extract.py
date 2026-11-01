"""验证提炼并发的顺序、限额和收尾，模型与 HTTP 均不访问网络。"""

import asyncio
from collections import Counter

import pytest

from tc_rag.domain import models as m
from tc_rag.domain.pipeline import Extraction


def extraction_requests(count):
    requests = []
    for index in range(count):
        text = f"第{index}项：最多300个，数量为整数。"
        block = m.DocumentBlock(
            block_id=f"b{index}",
            document_id="d",
            kind="paragraph",
            order=index,
            content=text,
            locator=m.Locator(paragraph_index=index),
        )
        chunk = m.Chunk(
            chunk_id=f"chunk{index}",
            text=text,
            token_estimate=30,
            fingerprint=f"fp{index}",
            sources=[m.ChunkSource(block_id=block.block_id, start_char=0, end_char=len(text))],
        )
        request = {
            "stage": "extract",
            "chunk": chunk.model_dump(mode="json"),
            "document_blocks": [block.model_dump(mode="json")],
        }
        requests.append((chunk, [block], request))
    return requests


@pytest.mark.asyncio
async def test_bounded_extract_keeps_source_order_when_completions_are_reversed(tmp_path):
    from test_pipeline import FixtureLLM, setup

    pipeline, _, _, _ = setup(tmp_path)
    pipeline.config.models.llm.concurrency = 2
    pipeline.config.runtime.cache_enabled = False
    entered = [asyncio.Event() for _ in range(5)]
    release = [asyncio.Event() for _ in range(5)]
    completed, events = [], []
    active = maximum = 0

    class ControlledLLM(FixtureLLM):
        async def generate(self, request, schema):
            nonlocal active, maximum
            index = request["document_blocks"][0]["order"]
            active += 1
            maximum = max(maximum, active)
            entered[index].set()
            try:
                await release[index].wait()
                completed.append(index)
                return await super().generate(request, schema)
            finally:
                active -= 1

    pipeline.llm = ControlledLLM()
    pipeline.progress = events.append
    task = asyncio.create_task(pipeline._extract_concurrently(extraction_requests(5), "run"))
    try:
        await asyncio.wait_for(entered[1].wait(), 2)
        assert entered[0].is_set() and not entered[2].is_set()
        # 保留第一批等待，后续四批先完成，验证返回结果仍按输入顺序排列。
        for index in range(1, 5):
            await asyncio.wait_for(entered[index].wait(), 2)
            release[index].set()
        release[0].set()
        results = await asyncio.wait_for(task, 2)
    finally:
        task.cancel()
        await asyncio.gather(task, return_exceptions=True)
    assert maximum == 2 and active == 0
    assert completed[0] == 1
    assert [result.evidence[0].block_id for result in results] == [f"b{i}" for i in range(5)]
    assert len({result.rules[0].rule_id for result in results}) == 5
    assert [e.completed for e in events if e.kind == "batch_completed"] == [1, 2, 3, 4, 5]


@pytest.mark.asyncio
@pytest.mark.parametrize("cancel_parent", [False, True])
async def test_extract_failure_or_cancel_drains_workers_without_starting_more(
    tmp_path,
    monkeypatch,
    cancel_parent,
):
    from test_pipeline import setup

    pipeline, _, _, _ = setup(tmp_path)
    pipeline.config.models.llm.concurrency = 3
    entered, finished, cancelled = [], [], []
    ready, fail = asyncio.Event(), asyncio.Event()

    async def controlled(chunk, blocks, request, run_id):
        index = blocks[0].order
        entered.append(index)
        if len(entered) == 3:
            ready.set()
        try:
            if index == 0 and not cancel_parent:
                await fail.wait()
                raise RuntimeError("fixture-extraction-failure")
            await asyncio.Event().wait()
            return Extraction()
        except asyncio.CancelledError:
            cancelled.append(index)
            raise
        finally:
            finished.append(index)

    monkeypatch.setattr(pipeline, "_extract_chunk", controlled)
    task = asyncio.create_task(pipeline._extract_concurrently(extraction_requests(6), "run"))
    try:
        await asyncio.wait_for(ready.wait(), 2)
        if cancel_parent:
            task.cancel()
            with pytest.raises(asyncio.CancelledError):
                await asyncio.wait_for(task, 2)
        else:
            fail.set()
            with pytest.raises(RuntimeError, match="fixture-extraction-failure"):
                await asyncio.wait_for(task, 2)
    finally:
        task.cancel()
        await asyncio.gather(task, return_exceptions=True)
    assert entered == [0, 1, 2]
    assert sorted(finished) == [0, 1, 2]
    assert sorted(cancelled) == ([0, 1, 2] if cancel_parent else [1, 2])


@pytest.mark.asyncio
async def test_planning_and_generation_are_bounded_and_keep_case_order(tmp_path):
    from docx import Document
    from test_pipeline import FixtureLLM, setup

    pipeline, doc, _, _ = setup(tmp_path)
    pipeline.config.models.llm.concurrency = 2
    pipeline.config.generation.batch_point_limit = 1
    document = Document()
    for title in ("第一模块", "第二模块"):
        document.add_heading(title, 1)
        document.add_paragraph("最多300个，数量为整数。")
    document.save(doc)
    active, maximum = Counter(), Counter()

    class ObservedLLM(FixtureLLM):
        async def generate(self, request, schema):
            stage = request["stage"]
            active[stage] += 1
            maximum[stage] = max(maximum[stage], active[stage])
            try:
                await asyncio.sleep(0)  # 确定性让出调度，检查在途批次而非比较运行时长。
                payload = dict(request)
                if stage == "extract":
                    # 规则证据必须引用含 300 与整数约束的正文，不能拿模块标题充当依据。
                    payload["document_blocks"] = [
                        next(b for b in request["document_blocks"] if b["kind"] == "paragraph")
                    ]
                result = await super().generate(payload, schema)
                if stage == "extract":
                    result.output.features[0].name = request["document_blocks"][0]["content"]
                return result
            finally:
                active[stage] -= 1

    pipeline.llm = ObservedLLM()
    result = await pipeline.generate(doc, "生成边界用例", tmp_path / "out", run_id="run_parallel")
    assert maximum["plan_points"] == maximum["generate"] == 2
    assert not any(active.values())
    payload = pipeline.read(result["run_id"], "cases")
    assert result["cases"] == 2
    assert [case["point_ids"][0] for case in payload["cases"]] == [
        point["point_id"] for point in payload["points"]
    ]
