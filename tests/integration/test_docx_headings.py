"""用虚构 DOCX 核对可选标题推断，避免排版启发式改变原文内容和定位。"""

from io import BytesIO

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
from PIL import Image

from tc_rag.infrastructure.parsing.docx import DocxParser
from tc_rag.infrastructure.storage.files import LocalArtifactStore


def parse_document(document, tmp_path, *, enabled=True):
    content = BytesIO()
    document.save(content)
    store = LocalArtifactStore(tmp_path / "managed")
    asset = store.put("fixtures/headings.docx", content.getvalue())
    return DocxParser(store, infer_headings=enabled).parse(asset)


def visual_title(document, text, size):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.bold = True
    return paragraph


@pytest.mark.parametrize("normal_declares_size", [False, True])
def test_inferred_headings_resolve_inherited_body_size(tmp_path, normal_declares_size):
    document = Document()
    if normal_declares_size:
        document.styles["Normal"].font.size = Pt(11)
    visual_title(document, "产品管理", 14)
    document.add_paragraph("保存后新增产品记录。")
    visual_title(document, "产品字段", 12)
    document.add_paragraph("产品名称为必填项。")
    result = parse_document(document, tmp_path)
    assert [block.kind for block in result.blocks] == [
        "heading", "paragraph", "heading", "paragraph"
    ]
    assert result.blocks[-1].heading_path == ["产品管理", "产品字段"]
    assert result.blocks[-1].locator.paragraph_index == 3
    assert any(item.severity == "info" and "排版推断" in item.description for item in result.issues)


def test_inference_resolves_character_and_paragraph_style_chains(tmp_path):
    document = Document()
    document.styles["Normal"].font.size = Pt(11)
    character_base = document.styles.add_style("VisualCharacterBase", WD_STYLE_TYPE.CHARACTER)
    character_base.font.size = Pt(14)
    character_base.font.bold = True
    character_child = document.styles.add_style("VisualCharacter", WD_STYLE_TYPE.CHARACTER)
    character_child.base_style = character_base
    title = document.add_paragraph("产品管理")
    title.runs[0].style = character_child
    paragraph_base = document.styles.add_style("VisualParagraphBase", WD_STYLE_TYPE.PARAGRAPH)
    paragraph_base.base_style = document.styles["Normal"]
    paragraph_base.font.size = Pt(12)
    paragraph_base.font.bold = True
    paragraph_child = document.styles.add_style("VisualParagraph", WD_STYLE_TYPE.PARAGRAPH)
    paragraph_child.base_style = paragraph_base
    document.add_paragraph("产品字段", style=paragraph_child)
    document.add_paragraph("产品名称为必填项。")
    result = parse_document(document, tmp_path)
    assert [block.kind for block in result.blocks] == ["heading", "heading", "paragraph"]
    assert result.blocks[-1].heading_path == ["产品管理", "产品字段"]


def test_inferred_subheadings_keep_explicit_parent_and_reset_at_next_heading(tmp_path):
    document = Document()
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading("产品管理", level=1)
    visual_title(document, "产品字段", 14)
    visual_title(document, "必填约束", 12)
    document.add_paragraph("产品名称为必填项。")
    document.add_heading("客户管理", level=1)
    visual_title(document, "客户字段", 14)
    document.add_paragraph("客户名称为必填项。")
    result = parse_document(document, tmp_path)
    assert result.blocks[3].heading_path == ["产品管理", "产品字段", "必填约束"]
    assert result.blocks[-1].heading_path == ["客户管理", "客户字段"]


def test_local_bold_emphasis_does_not_turn_entire_paragraph_into_heading(tmp_path):
    document = Document()
    document.styles["Normal"].font.size = Pt(11)
    paragraph = document.add_paragraph("当用户提交产品信息时，系统先校验")
    emphasis = paragraph.add_run("产品名称")
    emphasis.font.size = Pt(14)
    emphasis.bold = True
    paragraph.add_run("，校验失败则保留原页面供修改。")
    document.add_paragraph("校验成功后新增一条记录。")
    result = parse_document(document, tmp_path)
    assert [block.kind for block in result.blocks] == ["paragraph", "paragraph"]
    assert all(not block.heading_path for block in result.blocks)
    assert result.blocks[0].content == paragraph.text


@pytest.mark.parametrize("text", ["标题" * 41, "请输入所有必填字段后再保存。"])
def test_long_or_sentence_like_bold_text_stays_paragraph(tmp_path, text):
    document = Document()
    document.styles["Normal"].font.size = Pt(11)
    visual_title(document, text, 14)
    document.add_paragraph("正文。").runs[0].font.size = Pt(11)
    result = parse_document(document, tmp_path)
    assert result.blocks[0].kind == "paragraph"
    assert result.blocks[0].content == text


def test_explicit_nonbold_override_prevents_heading_inference(tmp_path):
    document = Document()
    style = document.styles.add_style("BoldParagraph", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles["Normal"]
    style.font.size = Pt(14)
    style.font.bold = True
    paragraph = document.add_paragraph("产品字段", style=style)
    paragraph.runs[0].bold = False
    result = parse_document(document, tmp_path)
    assert result.blocks[0].kind == "paragraph"


@pytest.mark.parametrize("enabled", [False, True])
def test_explicit_heading_is_preserved_with_optional_inference(tmp_path, enabled):
    document = Document()
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading("产品管理", level=1)
    visual_title(document, "产品字段", 14)
    document.add_paragraph("产品名称为必填项。")
    result = parse_document(document, tmp_path, enabled=enabled)
    assert result.blocks[0].kind == "heading"
    assert result.blocks[-1].heading_path == (
        ["产品管理", "产品字段"] if enabled else ["产品管理"]
    )
    assert result.blocks[1].locator.paragraph_index == 1


def test_heading_inference_defaults_to_disabled(tmp_path):
    document = Document()
    visual_title(document, "产品管理", 14)
    document.add_paragraph("产品名称为必填项。")
    content = BytesIO()
    document.save(content)
    store = LocalArtifactStore(tmp_path / "managed")
    result = DocxParser(store).parse(store.put("fixtures/headings.docx", content.getvalue()))
    assert all(block.kind == "paragraph" and not block.heading_path for block in result.blocks)
    assert not result.issues


def test_mixed_picture_paragraph_keeps_text_image_order_and_locator(tmp_path):
    document = Document()
    document.styles["Normal"].font.size = Pt(11)
    paragraph = visual_title(document, "图片前", 14)
    image = BytesIO()
    Image.new("RGB", (4, 4), "red").save(image, format="PNG")
    image.seek(0)
    paragraph.add_run().add_picture(image)
    tail = paragraph.add_run("图片后")
    tail.font.size = Pt(14)
    tail.bold = True
    result = parse_document(document, tmp_path)
    assert [block.kind for block in result.blocks] == ["paragraph", "image", "paragraph"]
    assert [block.content for block in result.blocks if block.kind == "paragraph"] == [
        "图片前", "图片后"
    ]
    assert all(block.locator.paragraph_index == 0 for block in result.blocks)
    assert all(not block.heading_path for block in result.blocks)
