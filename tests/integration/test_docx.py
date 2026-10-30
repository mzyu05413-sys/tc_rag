"""动态创建虚构DOCX，验证文件去重、正文顺序、表格和图片定位。"""

import json
from io import BytesIO

import pytest
from docx import Document
from PIL import Image

from tc_rag.application.services.ingest import ingest_document
from tc_rag.infrastructure.parsing.docx import DocxParser
from tc_rag.infrastructure.storage.files import LocalArtifactStore
from tc_rag.infrastructure.storage.sqlite import SqliteRepository


def sample(path):
    doc = Document()
    doc.add_heading("演示章节", level=1)
    doc.add_paragraph("保留否定词：不可小于0，单位元。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "规则"
    table.cell(1, 0).text = "演示金额"
    table.cell(1, 1).text = "大于0"
    p = doc.add_paragraph("图片前")
    image = BytesIO()
    Image.new("RGB", (4, 4), "red").save(image, format="PNG")
    image.seek(0)
    p.add_run().add_picture(image)
    p.add_run("图片后")
    doc.add_heading("子章节", level=2)
    doc.add_paragraph("结束")
    doc.save(path)


def test_ingest_parse_and_reuse(tmp_path):
    path = tmp_path / "demo.docx"
    sample(path)
    store = LocalArtifactStore(tmp_path / "managed")
    repo = SqliteRepository(tmp_path / "data.sqlite")
    first = ingest_document(path, store, repo)
    renamed = tmp_path / "renamed.docx"
    renamed.write_bytes(path.read_bytes())
    second = ingest_document(renamed, store, repo)
    assert not first.reused and second.reused
    assert first.document == second.document
    result = DocxParser(store).parse(first.asset)
    assert [b.kind for b in result.blocks] == [
        "heading",
        "paragraph",
        "table",
        "paragraph",
        "image",
        "paragraph",
        "heading",
        "paragraph",
    ]
    assert [b.order for b in result.blocks] == list(range(8))
    assert result.blocks[1].content == "保留否定词：不可小于0，单位元。"
    assert result.blocks[-1].heading_path == ["演示章节", "子章节"]
    assert json.loads(result.blocks[2].content)["rows"][1] == ["演示金额", "大于0"]
    assert result.blocks[4].locator.paragraph_index == 2
    assert store.read(result.assets[0]).startswith(b"\x89PNG")
    assert any(i.kind == "unparsed" and "图片" in i.description for i in result.issues)
    # 解析器返回数据，保存仍由调用方显式提交，以便后续统一控制事务边界。
    repo.save_many(
        [("assets", a.asset_id, a) for a in result.assets]
        + [("document_blocks", b.block_id, b) for b in result.blocks]
        + [("issues", i.issue_id, i) for i in result.issues]
    )


@pytest.mark.parametrize(
    "name,content",
    [("empty.docx", b""), ("bad.docx", b"not zip"), ("bad.pdf", b"not pdf"), ("bad.txt", b"text")],
)
def test_reject_bad_inputs(tmp_path, name, content):
    source = tmp_path / name
    source.write_bytes(content)
    store = LocalArtifactStore(tmp_path / "managed")
    repo = SqliteRepository(tmp_path / "data.sqlite")
    with pytest.raises(ValueError):
        ingest_document(source, store, repo)
    assert not (store.root / "assets").exists()


def test_size_missing_file_path_escape_and_checksum(tmp_path):
    store = LocalArtifactStore(tmp_path / "managed")
    repo = SqliteRepository(tmp_path / "data.sqlite")
    source = tmp_path / "demo.docx"
    with pytest.raises(ValueError):
        ingest_document(source, store, repo)
    sample(source)
    with pytest.raises(ValueError):
        ingest_document(source, store, repo, max_bytes=10)
    with pytest.raises(ValueError):
        store.put("../escape", b"demo")
    asset = store.put("safe/demo.bin", b"demo")
    with pytest.raises(ValueError):
        store.put("safe/demo.bin", b"different")
    (store.root / asset.relative_path).write_bytes(b"changed")
    with pytest.raises(ValueError, match="校验"):
        store.read(asset)


def test_merged_nested_and_empty_content_reported(tmp_path):
    path = tmp_path / "complex.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).add_table(rows=1, cols=1)
    doc.save(path)
    store = LocalArtifactStore(tmp_path / "managed")
    repo = SqliteRepository(tmp_path / "data.sqlite")
    result = DocxParser(store).parse(ingest_document(path, store, repo).asset)
    assert any("合并单元格" in issue.description for issue in result.issues)
    assert any("嵌套表格" in issue.description for issue in result.issues)


def test_size_and_bold_heading_levels(tmp_path):
    """未使用 Word Heading 样式、改用字号+加粗表达层级时也能建立 heading_path。"""
    from docx.shared import Pt

    path = tmp_path / "sized.docx"
    doc = Document()
    for text in ("正文一", "正文二", "正文三"):
        doc.add_paragraph(text).runs[0].font.size = Pt(11)
    first = doc.add_paragraph().add_run("产品列表")
    first.bold = True
    first.font.size = Pt(14)
    second = doc.add_paragraph().add_run("列表搜索")
    second.bold = True
    second.font.size = Pt(12)
    doc.add_paragraph("支持通过产品名称搜索").runs[0].font.size = Pt(11)
    doc.save(path)

    store = LocalArtifactStore(tmp_path / "managed")
    repo = SqliteRepository(tmp_path / "data.sqlite")
    result = DocxParser(store).parse(ingest_document(path, store, repo).asset)
    body = next(b for b in result.blocks if "支持通过产品名称搜索" in b.content)
    assert body.heading_path == ["产品列表", "列表搜索"]
    headings = [b.content for b in result.blocks if b.kind == "heading"]
    assert headings == ["产品列表", "列表搜索"]
