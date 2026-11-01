"""保留首张十一列模板，同时确保附表来源文本不会变成公式。"""

import json
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from tc_rag.infrastructure.exporters.reports import ReportExporter


def test_boundary_data_variants_are_visible_in_remarks(tmp_path):
    output = tmp_path / "data.xlsx"
    ReportExporter().export(
        {
            "cases": [
                {
                    "case_id": "c1",
                    "title": "学习承诺长度",
                    "status": "validated",
                    "data": {"data_variants": ["34", "35", "36"]},
                    "steps": [{"action": "输入", "expected": "保存"}],
                }
            ]
        },
        output,
    )
    sheet = load_workbook(output)["模版"]
    assert "[数据] 34、35、36" in sheet["H2"].value


def test_remarks_use_specific_case_issues(tmp_path):
    output = tmp_path / "remarks.xlsx"
    ReportExporter().export(
        {
            "cases": [
                {
                    "case_id": "c1",
                    "title": "t1",
                    "status": "needs_review",
                    "steps": [{"action": "a", "expected": "e"}],
                },
                {
                    "case_id": "c2",
                    "title": "t2",
                    "status": "needs_review",
                    "steps": [{"action": "a", "expected": "e"}],
                },
            ],
            "issues": [
                {
                    "issue_id": "i1",
                    "kind": "unsupported_expectation",
                    "description": "预期缺少依据",
                    "object_ids": ["c1"],
                }
            ],
        },
        output,
    )
    sheet = load_workbook(output)["模版"]
    assert "预期缺少依据" in sheet["H2"].value
    assert "详见问题清单" in sheet["H3"].value


def test_validated_case_uses_advisory_remark(tmp_path):
    output = tmp_path / "advisory.xlsx"
    ReportExporter().export(
        {
            "cases": [
                {
                    "case_id": "c1",
                    "title": "边界拒绝",
                    "status": "validated",
                    "steps": [{"action": "a", "expected": "e"}],
                }
            ],
            "issues": [
                {
                    "issue_id": "i1",
                    "kind": "unsupported_expectation",
                    "severity": "info",
                    "description": "预期为业务表达，未逐字出现在原文中",
                    "object_ids": ["c1"],
                }
            ],
        },
        output,
    )
    sheet = load_workbook(output)["模版"]
    # validated 且仅带 info 提示的用例，备注不再写 "[提示]"，保持干净可直接使用。
    remark = sheet["H2"].value or ""
    assert "[提示]" not in remark
    assert "[需核实]" not in remark


def test_quality_sheets_and_formula_text(tmp_path):
    output = tmp_path / "cases.xlsx"
    ReportExporter().export(
        {
            "cases": [],
            "rules": [{"rule_id": "r", "statement": "=1+1"}],
            "issues": [
                {"issue_id": "i", "kind": "unparsed", "description": "@source", "object_ids": []}
            ],
        },
        output,
    )
    book = load_workbook(output)
    assert book.sheetnames == [
        "模版",
        "阅读说明",
        "测试点",
        "规则",
        "覆盖矩阵",
        "问题清单",
        "未解析范围",
        "用例追溯",
        "质量说明",
    ]
    assert book["规则"]["B2"].value == "=1+1"
    assert book["规则"]["B2"].data_type == "s"
    assert book["未解析范围"].max_row == 2


def test_appendix_sheets_are_human_readable(tmp_path):
    output = tmp_path / "readable.xlsx"
    ReportExporter().export(
        {
            "cases": [{"case_id": "c1", "title": "提交订单", "steps": []}],
            "points": [
                {
                    "point_id": "p1",
                    "objective": "金额校验",
                    "category": "字段校验",
                    "technique": "boundary",
                    "rule_ids": ["r1"],
                    "status": "ready",
                }
            ],
            "rules": [
                {
                    "rule_id": "r1",
                    "kind": "field_constraint",
                    "statement": "金额大于0",
                    "evidence_ids": ["e1"],
                    "status": "supported",
                }
            ],
            "evidence": [
                {"evidence_id": "e1", "block_id": "b", "locator": {"page": 1}, "quote": "金额大于0"}
            ],
            "coverage": [
                {
                    "case_id": "c1",
                    "step_no": 1,
                    "rule_id": "r1",
                    "point_id": "p1",
                    "validation_status": "needs_review",
                }
            ],
            "issues": [
                {
                    "issue_id": "i1",
                    "kind": "coverage",
                    "severity": "warning",
                    "description": "测试点没有用例覆盖",
                    "object_ids": ["p1"],
                }
            ],
        },
        output,
    )
    book = load_workbook(output)
    try:
        assert list(book["测试点"].values)[1] == (
            "字段校验",
            "边界值",
            "金额校验",
            "金额大于0",
            "已就绪",
            "p1",
        )
        assert list(book["规则"].values)[1] == (
            "字段约束",
            "金额大于0",
            "金额大于0",
            "有依据",
            "r1",
        )
        assert list(book["覆盖矩阵"].values)[1] == (
            "提交订单",
            1,
            "金额大于0",
            "金额校验",
            "需人工核实",
        )
        assert list(book["问题清单"].values)[1] == (
            "覆盖缺口",
            "警告",
            "测试点没有用例覆盖",
            "测试点：金额校验",
            "具体问题或需求待明确",
            "待处理",
        )
    finally:
        book.close()


def test_custom_case_column_order_and_header(tmp_path):
    output = tmp_path / "custom.xlsx"
    ReportExporter([("priority", "级别"), ("title", "标题")]).export(
        {"cases": [{"title": "+source", "priority": "P1", "steps": []}]}, output
    )
    sheet = load_workbook(output).active
    assert list(sheet.values) == [("级别", "标题"), ("P1", "+source")]
    assert sheet["B2"].data_type == "s"


def test_illegal_controls_are_visible_in_excel_without_mutating_facts(tmp_path):
    # 覆盖 openpyxl 拒绝的全部 C0 控制字符，而不是只特判这次文档的 U+0001。
    codes = [*range(9), 11, 12, *range(14, 32)]
    original = "=中文" + "".join(chr(code) for code in codes) + "\t换行\n次行\r末行\\u0001"
    expected = "=中文" + "".join(f"\\u{code:04x}" for code in codes) + "\t换行\n次行\r末行\\u0001"
    payload = {
        "cases": [
            {
                "title": original,
                "external_id": "c",
                "status": "validated",
                "preconditions": [original],
                "remarks": original,
                "steps": [{"action": original, "expected": original}],
            }
        ],
        "points": [{"point_id": "p", "objective": original}],
        "rules": [{"rule_id": "r", "statement": original}],
        "coverage": [{"case_id": "c", "step_no": 1, "rule_id": "r", "point_id": "p"}],
        "issues": [
            {"issue_id": "i", "kind": "unparsed", "description": original, "object_ids": [original]}
        ],
    }
    before = deepcopy(payload)
    output = tmp_path / "controls.xlsx"
    exporter = ReportExporter()
    exporter.export(payload, output)
    book = load_workbook(output)
    try:
        sheet = book["模版"]
        assert sheet.max_column == 11 and sheet.max_row == 2
        assert sheet["A2"].value == expected and sheet["A2"].data_type == "s"
        assert sheet["C2"].value == expected and sheet["H2"].value == expected
        assert sheet["E2"].value == "[1] " + expected
        assert sheet["F2"].value == "[1] " + expected
        assert book["测试点"]["C2"].value == expected
        assert book["规则"]["B2"].value == expected
        assert book["问题清单"]["C2"].value == expected
        assert book["未解析范围"]["B2"].value == expected
        assert book["覆盖矩阵"]["B2"].value == 1
        assert book["问题清单"]["D2"].value == expected
        assert not any(cell.data_type == "f" for page in book for row in page for cell in row)
    finally:
        book.close()
    assert payload == before
    exporter.export(payload, tmp_path / "facts.json")
    assert json.loads((tmp_path / "facts.json").read_text(encoding="utf-8")) == before


def test_excel_controls_in_custom_headers_and_literal_escape(tmp_path):
    output = tmp_path / "custom-controls.xlsx"
    ReportExporter([("title", "标题\x01"), ("priority", "等级")]).export(
        {"cases": [{"title": r"字面量\u0001", "priority": None, "steps": []}]}, output
    )
    book = load_workbook(output)
    try:
        assert list(book["模版"].values) == [(r"标题\u0001", "等级"), (r"字面量\u0001", None)]
    finally:
        book.close()


@pytest.mark.parametrize(
    "description,severity",
    [
        ("预期为解释性断言，需人工核对语义", "warning"),
        ("预期为业务表达，未逐字出现在原文中，建议人工确认", "info"),
    ],
)
def test_trace_separates_generic_notices_case_risks_and_shared_source_issues(
    tmp_path, description, severity
):
    payload = {
        "cases": [
            {
                "case_id": case_id,
                "title": f"检查-{case_id}",
                "status": "needs_review",
                "remarks": "原始备注" if case_id == "c2" else None,
                "point_ids": ["p1"],
                "steps": [
                    {"action": "输入", "expected": "拒绝", "rule_ids": ["r1"],
                     "evidence_ids": ["e1"]},
                ],
            }
            for case_id in ("c1", "c2")
        ],
        "points": [{"point_id": "p1", "objective": "校验输入", "rule_ids": ["r1"]}],
        "rules": [{"rule_id": "r1", "statement": "允许范围有待确认", "evidence_ids": ["e1"]}],
        "evidence": [{"evidence_id": "e1", "quote": "范围由管理员配置"}],
        "issues": [
            {"issue_id": "i1", "kind": "unsupported_expectation", "severity": severity,
             "description": description, "object_ids": ["c1", "c2"]},
            {"issue_id": "i2", "kind": "unsupported_expectation", "severity": "warning",
             "description": "预期数值或证据范围缺乏直接依据", "object_ids": ["c2"]},
            {"issue_id": "i3", "kind": "gap", "severity": "warning",
             "description": "需求缺少允许范围", "object_ids": ["r1", "e1"]},
            {"issue_id": "i4", "kind": "other", "severity": "info",
             "description": "可读性提示", "object_ids": ["c1"]},
        ],
    }
    before = deepcopy(payload)
    output = tmp_path / "trace.xlsx"
    ReportExporter().export(payload, output)
    book = load_workbook(output)
    try:
        assert book["模版"]["H2"].value is None
        assert book["模版"]["H3"].value == (
            "原始备注\n[需核实] 预期数值或证据范围缺乏直接依据"
        )
        trace = book["用例追溯"]
        assert list(trace.values)[1] == (
            "c1", "检查-c1", "needs_review（需人工核实）", "p1", "校验输入",
            "r1", "允许范围有待确认", "e1", "范围由管理员配置", None, "需求缺少允许范围",
        )
        assert trace["J3"].value == "预期数值或证据范围缺乏直接依据"
        assert trace["K3"].value == "需求缺少允许范围"
        assert book["问题清单"].max_row == 5
        assert book["问题清单"]["E2"].value == "通用语义复核提醒"
        assert book["问题清单"]["E5"].value == "信息提示"
        summary = dict(list(book["质量说明"].values)[1:])
        assert summary["问题记录总数"] == 4
        assert summary["通用语义复核提醒记录数"] == 1
        assert summary["其他问题或提示记录数"] == 3
        assert summary["具体风险记录数（排除通用提醒与info）"] == 2
        assert summary["用例自身关联具体风险的数量"] == 1
        assert summary["关联需求或来源问题的用例数量"] == 2
        assert "不是缺陷或用例数" in summary["计数口径"]
    finally:
        book.close()
    assert payload == before


@pytest.mark.parametrize("status", ["validated", "future_status"])
def test_specific_risks_are_visible_without_rewriting_technical_status(tmp_path, status):
    payload = {
        "cases": [{"case_id": "c", "title": "检查", "status": status, "steps": []}],
        "issues": [
            {"issue_id": "i", "kind": "unsupported_expectation", "severity": "warning",
             "description": "明确的预期依据风险", "object_ids": ["c"]},
        ],
    }
    before = deepcopy(payload)
    output = tmp_path / "state.xlsx"
    ReportExporter().export(payload, output)
    book = load_workbook(output)
    try:
        assert "明确的预期依据风险" in book["模版"]["H2"].value
        assert book["用例追溯"]["C2"].value.startswith(status)
        assert book["用例追溯"]["J2"].value == "明确的预期依据风险"
    finally:
        book.close()
    assert payload == before


def test_new_trace_sheets_escape_controls_and_preserve_formula_text(tmp_path):
    payload = {
        "cases": [
            {"case_id": "=case", "title": "+title\x01", "status": "validated",
             "point_ids": ["@point"],
             "steps": [{"action": "a", "expected": "e", "rule_ids": ["-rule"],
                        "evidence_ids": ["=evidence"]}]},
        ],
        "points": [{"point_id": "@point", "objective": "=point\x01"}],
        "rules": [{"rule_id": "-rule", "statement": "=rule\x01"}],
        "evidence": [{"evidence_id": "=evidence", "quote": "@evidence\x01"}],
        "issues": [
            {"issue_id": "i", "kind": "gap", "description": "=source\x01",
             "severity": "warning", "object_ids": ["-rule"]},
        ],
    }
    before = deepcopy(payload)
    output = tmp_path / "trace-controls.xlsx"
    ReportExporter().export(payload, output)
    book = load_workbook(output)
    try:
        trace = book["用例追溯"]
        assert trace["A2"].value == "=case" and trace["A2"].data_type == "s"
        assert trace["B2"].value == r"+title\u0001"
        assert trace["E2"].value == r"=point\u0001"
        assert trace["G2"].value == r"=rule\u0001"
        assert trace["I2"].value == r"@evidence\u0001"
        assert trace["K2"].value == r"=source\u0001"
        assert not any(cell.data_type == "f" for page in book for row in page for cell in row)
        assert trace.freeze_panes == "A2" and trace.auto_filter.ref == trace.dimensions
    finally:
        book.close()
    assert payload == before
