"""导出可阅读报告及目标十一列 Excel；这些文件都是 SQLite 事实的视图。"""

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from tc_rag.domain.text import display_title

DEFAULT_COLUMNS = [
    ("title", "用例名称"),
    ("external_id", "ID"),
    ("preconditions", "前置条件"),
    ("module_path", "所属模块"),
    ("actions", "步骤描述"),
    ("expected", "预期结果"),
    ("tags", "标签"),
    ("remarks", "备注"),
    ("priority", "用例等级"),
    ("product_version", "CRM版本"),
    ("edit_mode", "编辑模式"),
]

TECHNIQUE_LABELS = {
    "equivalence": "等价类",
    "boundary": "边界值",
    "decision_table": "判定表",
    "state": "状态转换",
    "permission": "权限",
    "workflow": "流程",
    "error_behavior": "异常行为",
    "other": "其他",
}
POINT_STATUS_LABELS = {
    "ready": "已就绪",
    "needs_context": "待补充上下文",
    "needs_clarification": "待澄清",
    "not_applicable": "不适用",
}
RULE_KIND_LABELS = {
    "field_constraint": "字段约束",
    "business_condition": "业务条件",
    "permission": "权限",
    "state_transition": "状态流转",
    "error_behavior": "异常行为",
    "workflow": "流程",
    "display": "界面展示",
    "other": "其他",
}
RULE_STATUS_LABELS = {
    "supported": "有依据",
    "ambiguous": "含义不明",
    "conflicting": "存在冲突",
    "needs_review": "待审核",
}
VALIDATION_LABELS = {
    "pending": "待校验",
    "validated": "已校验",
    "invalid": "无效",
    "needs_review": "需人工核实",
}
ISSUE_KIND_LABELS = {
    "gap": "需求缺口",
    "conflict": "需求冲突",
    "validation": "校验问题",
    "unsupported": "无依据",
    "unsupported_expectation": "无依据预期",
    "duplicate": "疑似重复",
    "coverage": "覆盖缺口",
    "unparsed": "未解析",
    "other": "其他",
}
SEVERITY_LABELS = {"info": "提示", "warning": "警告", "error": "错误"}


def _semantic_notice(item: dict[str, Any]) -> bool:
    """只识别既有通用词面提醒，不将其他无依据预期降级成提示。"""
    return item.get("kind") == "unsupported_expectation" and item.get("description") in {
        "预期为解释性断言，需人工核对语义",
        "预期为业务表达，未逐字出现在原文中，建议人工确认",
    }


def _specific_risk(item: dict[str, Any]) -> bool:
    """旧记录未写严重程度时沿用 warning；通用提醒和 info 不当作具体风险。"""
    return not _semantic_notice(item) and (item.get("severity") or "warning") != "info"


def _issue_category(item: dict[str, Any]) -> str:
    if _semantic_notice(item):
        return "通用语义复核提醒"
    return "信息提示" if item.get("severity") == "info" else "具体问题或需求待明确"


def _data_text(data: Any) -> str:
    """把测试数据变体放进主表备注；十一列模板不新增列。"""
    if not isinstance(data, dict) or not data:
        return ""
    variants = data.get("data_variants")
    if isinstance(variants, list) and variants:
        return "、".join(str(item) for item in variants)
    return json.dumps(data, ensure_ascii=False, separators=(",", ":"))


def _excel_value(value: Any) -> Any:
    """仅改变 Excel 展示：控制字符可见转义，原始事实及换行、制表符不变。"""
    if isinstance(value, str):
        # 必须在 append 前处理；openpyxl 在赋值时就会拒绝 XML 不允许的控制字符。
        # 不直接删除，保留字符码供用户追溯；已有字面量转义不会被再次转换。
        return ILLEGAL_CHARACTERS_RE.sub(lambda match: f"\\u{ord(match[0]):04x}", value)
    return value


class ReportExporter:
    def __init__(self, case_columns: list[tuple[str, str]] | None = None):
        self.case_columns = case_columns or DEFAULT_COLUMNS

    def export(self, payload: dict[str, Any], destination: Path) -> Path:
        destination.parent.mkdir(parents=True, exist_ok=True)
        suffix = destination.suffix.lower()
        if suffix == ".json":
            destination.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
            )
        elif suffix == ".md":
            lines = [
                "# 测试设计运行报告",
                "",
                "自动检查不代表人工批准；语义、覆盖率及业务预期需要审核。",
                "",
            ]
            labels = {
                "rules": "需求规则",
                "points": "测试点",
                "cases": "测试用例",
                "evidence": "原文证据",
                "issues": "待审核问题",
                "coverage": "覆盖关联",
                "features": "功能",
                "quality": "质量统计",
                "document": "文档来源",
            }
            for section, values in payload.items():
                if section in {"schema_version", "tree", "entities"}:
                    continue
                lines.extend([f"## {labels.get(section, section)}", ""])
                if isinstance(values, list):
                    for index, item in enumerate(values, 1):
                        if not isinstance(item, dict):
                            lines.extend([f"- {item}", ""])
                            continue
                        title = next(
                            (
                                item[k]
                                for k in ("title", "objective", "statement", "name", "description")
                                if item.get(k)
                            ),
                            f"记录 {index}",
                        )
                        lines.extend([f"### {index}. {title}", ""])
                        for key, value in item.items():
                            if key == "schema_version":
                                continue
                            if key == "steps":
                                for number, step in enumerate(value, 1):
                                    lines.extend(
                                        [
                                            f"{number}. 操作：{step['action']}",
                                            f"   预期：{step['expected']}",
                                            "",
                                        ]
                                    )
                            else:
                                rendered = (
                                    json.dumps(value, ensure_ascii=False)
                                    if isinstance(value, (dict, list))
                                    else str(value)
                                )
                                lines.extend([f"- {key}：{rendered}", ""])
                elif isinstance(values, dict):
                    lines.extend([f"- {key}：{value}\n" for key, value in values.items()])
                else:
                    lines.extend([str(values), ""])
            destination.write_text("\n".join(lines), encoding="utf-8")
        elif suffix == ".xlsx":
            self._excel(payload.get("cases", []), destination, payload, self.case_columns)
        else:
            raise ValueError("导出仅支持 json/md/xlsx")
        return destination

    @staticmethod
    def _excel(
        cases: list[dict[str, Any]],
        destination: Path,
        payload: dict[str, Any] | None = None,
        columns: list[tuple[str, str]] | None = None,
    ) -> None:
        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.title = "模版"
        columns = columns or DEFAULT_COLUMNS
        payload = payload or {}
        issues = payload.get("issues", [])
        by_object: dict[str, list[dict[str, Any]]] = {}
        for item in issues:
            for object_id in item.get("object_ids", []) or []:
                by_object.setdefault(str(object_id), []).append(item)
        sheet.append([_excel_value(header) for _, header in columns])
        for case in cases:
            steps = case["steps"]
            remarks = case.get("remarks") or ""
            own_issues = by_object.get(str(case.get("case_id")), [])
            blocking = list(
                dict.fromkeys(
                    str(item["description"])
                    for item in own_issues
                    if _specific_risk(item) and item.get("description")
                )
            )
            if blocking:
                remarks += "\n[需核实] " + "；".join(blocking[:2])
                if len(blocking) > 2:
                    remarks += f"；另有 {len(blocking) - 2} 条，详见用例追溯"
            elif case.get("status") != "validated" and not own_issues:
                # 未记录原因时保留既有状态兜底，不能用空备注掩盖未通过状态。
                remarks += "\n[需核实] 自动校验未通过，详见问题清单"
            data_text = _data_text(case.get("data"))
            if data_text:
                remarks += "\n[数据] " + data_text
            # 新旧通用语义提醒都保留在问题清单/质量说明中，不逐用例重复。
            # 具体来源问题在追溯表说明；导出仅增加视图，不改状态和 Issue 原始事实。
            values = {
                **case,
                "title": display_title(str(case.get("title", ""))),
                "preconditions": "\n".join(case.get("preconditions", [])),
                "actions": "\n".join(f"[{i}] {s['action']}" for i, s in enumerate(steps, 1)),
                "expected": "\n".join(f"[{i}] {s['expected']}" for i, s in enumerate(steps, 1)),
                "tags": ",".join(case.get("tags", [])),
                "remarks": remarks.strip(),
                "edit_mode": case.get("edit_mode", "STEP"),
            }
            sheet.append([_excel_value(values.get(field)) for field, _ in columns])
            # 强制文本，避免把需求中的 =、+、-、@ 解释成电子表格公式。
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
        for index, (field, _) in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = (
                55 if field in {"actions", "expected"} else 26
            )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        point_map = {item.get("point_id"): item for item in payload.get("points", [])}
        rule_map = {item.get("rule_id"): item for item in payload.get("rules", [])}
        evidence_map = {item.get("evidence_id"): item for item in payload.get("evidence", [])}
        case_map = {item.get("case_id"): item for item in payload.get("cases", [])}

        def short(value: Any, limit: int = 80) -> str:
            text = "" if value is None else str(value)
            return text if len(text) <= limit else text[: limit - 1] + "…"

        def rule_text(rule_id: str) -> str:
            rule = rule_map.get(rule_id)
            return short(rule.get("statement"), 60) if rule else str(rule_id)

        def point_text(point_id: str) -> str:
            point = point_map.get(point_id)
            return short(point.get("objective"), 60) if point else str(point_id)

        def object_text(object_id: str) -> str:
            if object_id in case_map:
                return "用例：" + short(case_map[object_id].get("title"), 40)
            if object_id in point_map:
                return "测试点：" + point_text(object_id)
            if object_id in rule_map:
                return "规则：" + rule_text(object_id)
            evidence_item = evidence_map.get(object_id)
            if evidence_item:
                return "证据：" + short(evidence_item.get("quote"), 40)
            return str(object_id)

        trace_rows: list[list[Any]] = []
        related_count = own_risk_count = 0
        for case in cases:
            point_ids = list(dict.fromkeys(case.get("point_ids") or []))
            rule_ids = sorted({r for step in case["steps"] for r in step.get("rule_ids", [])})
            evidence_ids = sorted(
                {e for step in case["steps"] for e in step.get("evidence_ids", [])}
            )
            source_ids = {*point_ids, *rule_ids, *evidence_ids}
            own = by_object.get(str(case.get("case_id")), [])
            own_risks = list(
                dict.fromkeys(
                    str(item["description"])
                    for item in own
                    if _specific_risk(item) and item.get("description")
                )
            )
            related = list(
                dict.fromkeys(
                    str(item["description"])
                    for object_id in sorted(source_ids)
                    for item in by_object.get(object_id, [])
                    if not _semantic_notice(item) and item.get("description")
                )
            )
            related_count += bool(related)
            own_risk_count += bool(own_risks)
            status = str(case.get("status") or "未记录")
            label = VALIDATION_LABELS.get(status)
            trace_rows.append(
                [
                    case.get("case_id"),
                    display_title(str(case.get("title", ""))),
                    f"{status}（{label}）" if label else status,
                    "\n".join(point_ids),
                    "\n".join(point_text(p) for p in point_ids),
                    "\n".join(rule_ids),
                    "\n".join(rule_text(r) for r in rule_ids),
                    "\n".join(evidence_ids),
                    "\n".join(
                        short(evidence_map.get(e, {}).get("quote", e), 120) for e in evidence_ids
                    ),
                    "\n".join(own_risks),
                    "\n".join(related),
                ]
            )

        sheets: list[tuple[str, list[str], list[list[Any]]]] = [
            (
                "阅读说明",
                ["工作表", "用途"],
                [
                    ["模版", "测试用例主表，步骤与预期一一对应，可导入测试管理平台"],
                    ["测试点", "按分类和方法整理的测试点清单"],
                    ["规则", "从需求原文提炼的规则及证据摘要"],
                    ["覆盖矩阵", "用例步骤与规则、测试点的对应关系"],
                    ["问题清单", "自动校验发现的缺口、冲突和待确认项"],
                    ["未解析范围", "未能解析的图片或页面，需要人工补充"],
                    [
                        "用例追溯",
                        "按用例查看测试点、规则、证据摘要和关联问题，保留原始 ID 与技术状态",
                    ],
                    ["质量说明", "区分通用提醒、具体风险及来源问题，说明数量口径与自动校验边界"],
                ],
            ),
            (
                "测试点",
                ["分类", "测试方法", "测试目标", "关联规则", "状态", "参考ID"],
                [
                    [
                        point.get("category", ""),
                        TECHNIQUE_LABELS.get(
                            point.get("technique"), point.get("technique", "")
                        ),
                        point.get("objective", ""),
                        "；".join(rule_text(rid) for rid in (point.get("rule_ids") or [])[:2]),
                        POINT_STATUS_LABELS.get(point.get("status"), point.get("status", "")),
                        point.get("point_id", ""),
                    ]
                    for point in payload.get("points", [])
                ],
            ),
            (
                "规则",
                ["规则类型", "规则内容", "证据摘要", "状态", "参考ID"],
                [
                    [
                        RULE_KIND_LABELS.get(rule.get("kind"), rule.get("kind", "")),
                        rule.get("statement", ""),
                        short(
                            "；".join(
                                short(evidence_map[eid].get("quote"), 50)
                                for eid in (rule.get("evidence_ids") or [])[:1]
                                if eid in evidence_map
                            ),
                            120,
                        ),
                        RULE_STATUS_LABELS.get(rule.get("status"), rule.get("status", "")),
                        rule.get("rule_id", ""),
                    ]
                    for rule in payload.get("rules", [])
                ],
            ),
            (
                "覆盖矩阵",
                ["用例名称", "步骤", "规则内容", "测试目标", "校验状态"],
                [
                    [
                        short(case_map.get(link.get("case_id"), {}).get("title"), 50),
                        link.get("step_no", ""),
                        rule_text(link.get("rule_id", "")),
                        point_text(link.get("point_id", "")),
                        VALIDATION_LABELS.get(
                            link.get("validation_status"),
                            link.get("validation_status", ""),
                        ),
                    ]
                    for link in payload.get("coverage", [])
                ],
            ),
            (
                "问题清单",
                ["类型", "严重程度", "说明", "涉及对象", "问题性质", "状态"],
                [
                    [
                        ISSUE_KIND_LABELS.get(item.get("kind"), item.get("kind", "")),
                        SEVERITY_LABELS.get(item.get("severity"), item.get("severity", "")),
                        item.get("description", ""),
                        short(
                            "；".join(
                                object_text(oid) for oid in (item.get("object_ids") or [])[:2]
                            ),
                            120,
                        ),
                        _issue_category(item),
                        {"open": "待处理", "resolved": "已解决", "ignored": "已忽略"}.get(
                            item.get("status"), "待处理"
                        ),
                    ]
                    for item in payload.get("issues", [])
                ],
            ),
            (
                "未解析范围",
                ["说明", "涉及对象"],
                [
                    [
                        item.get("description", ""),
                        short(
                            "；".join(
                                object_text(oid) for oid in (item.get("object_ids") or [])[:2]
                            ),
                            120,
                        ),
                    ]
                    for item in payload.get("issues", [])
                    if item.get("kind") == "unparsed"
                ],
            ),
            (
                "用例追溯",
                [
                    "用例ID", "用例名称", "技术状态", "测试点ID", "测试目标摘要",
                    "规则ID", "规则摘要", "证据ID", "证据摘要", "用例自身具体风险",
                    "关联需求或来源问题（不含通用语义提醒）",
                ],
                trace_rows,
            ),
            (
                "质量说明",
                ["项目", "说明 / 数量"],
                [
                    ["用例数量", len(cases)],
                    ["问题记录总数", len(issues)],
                    ["通用语义复核提醒记录数", sum(_semantic_notice(item) for item in issues)],
                    ["其他问题或提示记录数", sum(not _semantic_notice(item) for item in issues)],
                    [
                        "具体风险记录数（排除通用提醒与info）",
                        sum(_specific_risk(item) for item in issues),
                    ],
                    ["用例自身关联具体风险的数量", own_risk_count],
                    ["关联需求或来源问题的用例数量", related_count],
                    ["主表备注", "仅显示用例自身具体风险；原因未记录的未通过状态保留兜底提示。"],
                    ["通用提醒", "预期非逐字原文不等于错误；通用语义提醒不计作已证实缺陷。"],
                    ["来源问题", "共享的需求、测试点和证据问题见用例追溯；可能关联多条用例。"],
                    [
                        "计数口径",
                        "一条记录可能关联多对象，一条用例也可能关联多记录；记录数不是缺陷或用例数。",
                    ],
                    [
                        "状态边界",
                        "原始技术状态与问题完整保留；空备注不代表人工批准或完整语义正确。",
                    ],
                ],
            ),
        ]
        for title, headers, rows in sheets:
            appendix = book.create_sheet(title)
            appendix.append(headers)
            for row in rows:
                appendix.append([_excel_value(value) for value in row])
                for cell in appendix[appendix.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            appendix.freeze_panes = "A2"
            appendix.auto_filter.ref = appendix.dimensions
            for cell in appendix[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="305496")
            for index, header in enumerate(headers, 1):
                appendix.column_dimensions[get_column_letter(index)].width = (
                    65 if any(word in header for word in ("摘要", "说明", "问题", "风险")) else 28
                )
        book.save(destination)
