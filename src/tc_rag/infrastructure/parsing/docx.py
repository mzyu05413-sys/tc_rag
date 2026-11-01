"""按 DOCX 主文档顺序解析文本、表格与图片，显式报告暂不支持的内容。"""

import hashlib
import json
import re
from collections import Counter
from io import BytesIO
from typing import Any, Literal

from docx import Document as OpenDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from tc_rag.application.ports.contracts import ArtifactStore, ParsedDocument
from tc_rag.domain.models import Asset, DocumentBlock, Issue, Locator


class DocxParser:
    def __init__(self, store: ArtifactStore, infer_headings: bool = False) -> None:
        self.store = store
        # 标题推断会改变 heading_path 和后续分块；默认只认 Word 明确的标题样式。
        self.infer_headings = infer_headings

    def parse(self, asset: Asset) -> ParsedDocument:
        """DOCX 不提供可靠排版页码，使用段落/表格/单元格索引定位。"""
        document = OpenDocument(BytesIO(self.store.read(asset)))
        result = ParsedDocument(blocks=[])
        document_id = "doc_" + asset.checksum
        headings: dict[int, str] = {}
        inferred_levels = self._visual_headings(document) if self.infer_headings else {}
        if inferred_levels:
            result.issues.append(
                Issue(
                    issue_id=f"{document_id}_issue_0",
                    kind="validation",
                    severity="info",
                    object_ids=[document_id],
                    description=(
                        f"按字号与整段加粗推断 {len(inferred_levels)} 个标题；"
                        "章节层级为排版推断，请核对原文"
                    ),
                )
            )

        def add(
            kind: Literal["heading", "paragraph", "table", "image", "other"],
            text: str,
            locator: Locator,
        ) -> None:
            order = len(result.blocks)
            result.blocks.append(
                DocumentBlock(
                    block_id=f"{document_id}_block_{order}",
                    document_id=document_id,
                    kind=kind,
                    heading_path=[headings[level] for level in sorted(headings)],
                    order=order,
                    content=text,
                    locator=locator,
                )
            )

        def issue(description: str, extra_ids: list[str] | None = None) -> None:
            result.issues.append(
                Issue(
                    issue_id=f"{document_id}_issue_{len(result.issues)}",
                    kind="unparsed",
                    severity="warning",
                    object_ids=[document_id, *(extra_ids or [])],
                    description=description,
                )
            )

        def paragraph(element: Any, locator: Locator, include_text: bool = True) -> None:
            # 遍历XML内容而不是只用paragraph.runs，保留超链接中的文字和图片锚点顺序。
            pending: list[str] = []

            def flush() -> None:
                if pending and include_text:
                    add("paragraph", "".join(pending), locator)
                pending.clear()

            for node in element.iter():
                if node.tag == qn("w:t"):
                    pending.append(node.text or "")
                elif node.tag == qn("w:tab"):
                    pending.append("\t")
                elif node.tag == qn("w:br"):
                    pending.append("\n")
                elif node.tag == qn("a:blip"):
                    flush()
                    relationship = node.get(qn("r:embed"))
                    if not relationship or relationship not in document.part.related_parts:
                        issue(f"位置 {locator.model_dump_json()} 的外链或缺失图片未解析")
                        continue
                    part = document.part.related_parts[relationship]
                    content = part.blob
                    checksum = hashlib.sha256(content).hexdigest()
                    suffix = str(part.partname).rsplit(".", 1)[-1]
                    image = self.store.put(
                        f"assets/{asset.checksum}/images/{checksum}.{suffix}", content
                    )
                    result.assets.append(image)
                    add("image", json.dumps(image.model_dump(), ensure_ascii=False), locator)
                    issue(
                        f"图片 {image.asset_id} 仅提取与定位，图片文字尚未解析",
                        [image.asset_id],
                    )
                elif node.tag in {qn("w:txbxContent"), qn("w:object"), qn("w:pict")}:
                    issue(
                        f"位置 {locator.model_dump_json()} 含文本框或嵌入对象，布局/内容可能不完整"
                    )
            flush()

        paragraph_index = table_index = 0
        for child in document.element.body:
            if child.tag == qn("w:p"):
                value = Paragraph(child, document)
                level = self._explicit_level(value)
                if level is None:
                    level = inferred_levels.get(paragraph_index)
                locator = Locator(paragraph_index=paragraph_index)
                if level is not None:
                    headings = {key: text for key, text in headings.items() if key < level}
                    headings[level] = value.text
                    add("heading", value.text, locator)
                    paragraph(child, locator, include_text=False)
                else:
                    paragraph(child, locator)
                paragraph_index += 1
            elif child.tag == qn("w:tbl"):
                table = Table(child, document)
                # 保留二维行列而不扁平拼接；合并单元格的重复值不当成独立业务事实。
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                add(
                    "table",
                    json.dumps({"schema_version": "1.0", "rows": rows}, ensure_ascii=False),
                    Locator(table_index=table_index),
                )
                if child.xpath(".//w:gridSpan | .//w:vMerge"):
                    issue(f"表格 {table_index} 含合并单元格，行列语义需要人工复核")
                for row_index, row in enumerate(table.rows):
                    seen = set()
                    for column_index, cell in enumerate(row.cells):
                        if cell._tc in seen:
                            continue
                        seen.add(cell._tc)
                        if cell.tables:
                            issue(f"表格 {table_index} 单元格含嵌套表格，未展开")
                        for p in cell.paragraphs:
                            paragraph(
                                p._p,
                                Locator(
                                    table_index=table_index,
                                    row_index=row_index,
                                    column_index=column_index,
                                ),
                                include_text=False,
                            )
                table_index += 1
            elif child.tag != qn("w:sectPr"):
                issue(f"主文档含未支持元素 {child.tag}，未自动展开")
        for part in document.part.package.parts:
            name = str(part.partname)
            if any(key in name for key in ("/header", "/footer", "/footnotes", "/endnotes")):
                issue(f"附属内容 {name} 尚未解析")
        if not result.blocks:
            issue("主文档没有可解析正文")
        return result

    @staticmethod
    def _explicit_level(value: Paragraph) -> int | None:
        style = value.style.name if value.style else ""
        match = re.fullmatch(r"Heading (\d+)", style)
        return int(match.group(1)) if match else None

    @staticmethod
    def _style_font(style: Any, field: str) -> Any:
        """读取样式继承链，None 表示未声明，False 必须保留为明确的非粗体。"""
        seen: set[str] = set()
        while style is not None and style.style_id not in seen:
            seen.add(style.style_id)
            value = getattr(style.font, field)
            if value is not None:
                return value
            style = style.base_style
        return None

    @classmethod
    def _default_font(cls, document: Any, field: str) -> Any:
        try:
            normal = document.styles["Normal"]
        except KeyError:
            normal = None
        value = cls._style_font(normal, field)
        if value is not None:
            return value.pt if field == "size" else value
        # 常见 DOCX 的 Normal 本身不写字号，而是在 docDefaults 中声明 11pt。
        tag = "sz" if field == "size" else "b"
        nodes = document.styles.element.xpath(f"w:docDefaults/w:rPrDefault/w:rPr/w:{tag}")
        if not nodes:
            return None if field == "size" else False
        raw = nodes[0].get(qn("w:val"))
        if field == "bold":
            return raw not in {"0", "false", "off"}
        try:
            return float(raw) / 2
        except (TypeError, ValueError):
            # 排版信息未知时不据此推断标题，原文仍按普通段落完整保留。
            return None

    @classmethod
    def _effective_font(cls, run: Any, paragraph: Paragraph, field: str, default: Any) -> Any:
        value = getattr(run.font, field)
        if value is None:
            value = cls._style_font(run.style, field)
        if value is None:
            value = cls._style_font(paragraph.style, field)
        if value is None:
            return default
        return float(value.pt) if field == "size" else value

    @classmethod
    def _visual_headings(cls, document: Any) -> dict[int, int]:
        """保守补认视觉标题；不以局部强调替代标题，也不覆盖显式章节的父层级。"""
        paragraphs = [
            Paragraph(child, document)
            for child in document.element.body
            if child.tag == qn("w:p")
        ]
        default_size = cls._default_font(document, "size")
        default_bold = cls._default_font(document, "bold")
        body_sizes: Counter[float] = Counter()
        candidates: dict[int, float] = {}
        explicit: dict[int, int] = {}
        for index, paragraph in enumerate(paragraphs):
            level = cls._explicit_level(paragraph)
            if level is not None:
                explicit[index] = level
                continue
            formats: list[tuple[float | None, bool]] = []
            run_text: list[str] = []
            for run in paragraph.runs:
                visible = re.sub(r"\s+", "", run.text)
                if not visible:
                    continue
                size = cls._effective_font(run, paragraph, "size", default_size)
                bold = bool(cls._effective_font(run, paragraph, "bold", default_bold))
                formats.append((size, bold))
                run_text.append(visible)
                if not bold and size is not None:
                    body_sizes[size] += len(visible)
            text = paragraph.text.strip()
            sizes = {size for size, _ in formats}
            # 整段格式一致才参与推断；超链接、图文混排及长句保持原有解析方式。
            if (
                formats
                and len(re.sub(r"\s+", "", text)) <= 80
                and "\n" not in text
                and "\r" not in text
                and not text.endswith(("。", "！", "？", ";", "；", "!", "?"))
                and all(bold for _, bold in formats)
                and len(sizes) == 1
                and None not in sizes
                and "".join(run_text) == re.sub(r"\s+", "", text)
                and not paragraph._p.xpath(".//w:drawing | .//w:object | .//w:pict")
            ):
                size = formats[0][0]
                if size is not None:
                    candidates[index] = size
        body_size = body_sizes.most_common(1)[0][0] if body_sizes else default_size
        if body_size is None:
            return {}
        candidates = {index: size for index, size in candidates.items() if size > body_size + 0.05}

        # 每个显式标题区间独立排字号档位；视觉子标题最浅也在显式标题之下。
        # 不凭字号把 Heading 1 的子节点升成另一棵一级章节。
        anchor = -1
        sections: dict[int, list[int]] = {}
        for index in range(len(paragraphs)):
            if index in explicit:
                anchor = index
            elif index in candidates:
                sections.setdefault(anchor, []).append(index)
        inferred: dict[int, int] = {}
        for anchor, indices in sections.items():
            heading_sizes = sorted({candidates[index] for index in indices}, reverse=True)
            for index in indices:
                inferred[index] = (
                    explicit.get(anchor, 0) + heading_sizes.index(candidates[index]) + 1
                )
        return inferred
